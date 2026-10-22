from datetime import date, datetime, time, timedelta, timezone
from math import atan2, cos, radians, sin, sqrt
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

import jwt
from fastapi import Depends, FastAPI, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pwdlib import PasswordHash
from sqlalchemy import and_, or_, select, inspect, text, delete, func, update
from sqlalchemy.orm import Session
import base64
import re
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import EmailStr, TypeAdapter
from openpyxl.comments import Comment
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from .config import settings
from .database import Base, SessionLocal, engine
from .models import AttendanceRecord, QrSession, User, VacationRequest,PermissionRequest,UserPreferences,CompanySettings,Notification,SpaceBooking,Space,Holiday,EmployeeProfile,IncidenceCode,Post, PostResponse
from .schemas import (
    AdminVacationCreate,
    AttendanceAdminOut,
    AttendanceInput,
    AttendanceOut,
    AttendanceTodayOut,
    LoginInput,
    ManualAttendanceInput,
    QrAttendanceInput,
    TokenOut,
    UserCreate,
    UserOut,
    VacationRequestAdminOut,
    VacationRequestCreate,
    VacationRequestOut,
    VacationReviewInput,
    QrOut,
    PermissionRequestAdminOut,
    PermissionRequestCreate,
    PermissionRequestOut,
    PermissionReviewInput,
    VacationBalanceOut,
    count_workdays,
    ProfileUpdate,
    ProfileOut,
    PasswordChange,
    NotificationSettings,
    TodayAbsenceOut,
    TodayOverviewOut,
    TodayEntryOut,
    UserCreatedOut,
    UserUpdate,
    AvatarInput,
    BrandingOut,
    BrandingUpdate,
    NotificationOut,
    NotificationsOut,
    SpaceIn,
    SpaceOut,
    BookingIn,
    BookingOut,
    to_minutes,
    to_hhmm,
    AttendanceEdit,
    ManualAttendanceInput,
    HolidayIn,
    HolidayOut,
    HHMM,
    ScheduleOut,
    PersonOut,
    ProfileUpdateIn,
    ImportIn,
    ImportResult,
    ImportRowResult,
    ImportCredential,
    IncidenceCodeIn,
    IncidenceCodeOut,
    ReportRequest,
    ReportPersonOut,
    ReportPreviewOut,
    PostIn,
    PostOut,
    PostAuthorOut,
    PostsMetaOut,
    RespondIn,
    ResponsePersonOut,
    CalendarEventOut
)


app = FastAPI(
    title="API de asistencia móvil",
    version="0.2.0",
)


# =========================================================
# CONFIGURACIÓN
# =========================================================

LOCAL_TIMEZONE = ZoneInfo(settings.timezone)

password_hash = PasswordHash.recommended()

JWT_ALGORITHM = "HS256"

VACATION_BASE_DAYS = 12
VACATION_YEARLY_INCREMENT = 2

AVATAR_FORMATS = {
    "data:image/jpeg;base64,": b"\xff\xd8\xff",
    "data:image/png;base64,": b"\x89PNG",
    "data:image/webp;base64,": b"RIFF",
}

MAX_AVATAR_BYTES = 300_000

PREF_BY_KIND = {
    "missing_exit": "notify_attendance",
    "attendance_fix": "notify_attendance",
    "vacation": "notify_vacations",
    "permission": "notify_permissions",
    "weekly": "notify_weekly",
    "company": "notify_company",
    "meeting_reminder": "notify_company"
}


# =========================================================
# CORS
# =========================================================

cors_origins = [
    origin.strip()
    for origin in settings.cors_origins.split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# UTILIDADES
# =========================================================


def distance_meters(
    latitude1: float,
    longitude1: float,
    latitude2: float,
    longitude2: float,
) -> float:
    """
    Calcula la distancia aproximada entre dos coordenadas
    utilizando la fórmula de Haversine.
    """

    earth_radius = 6_371_000

    lat1 = radians(latitude1)
    lat2 = radians(latitude2)

    delta_lat = radians(latitude2 - latitude1)
    delta_lon = radians(longitude2 - longitude1)

    a = (
        sin(delta_lat / 2) ** 2
        + cos(lat1)
        * cos(lat2)
        * sin(delta_lon / 2) ** 2
    )

    c = 2 * atan2(sqrt(a), sqrt(1 - a))

    return earth_radius * c


def get_local_now() -> datetime:
    """
    Obtiene la fecha/hora actual de la empresa.
    """

    return datetime.now(timezone.utc).astimezone(
        LOCAL_TIMEZONE
    )


def get_work_date(
    value: datetime | None = None,
) -> date:
    """
    Obtiene la fecha laboral utilizando el timezone
    configurado para la empresa.
    """

    if value is None:
        return get_local_now().date()

    return value.astimezone(LOCAL_TIMEZONE).date()


def normalize_datetime(value: datetime) -> datetime:
    """
    Convierte una fecha/hora recibida sin timezone
    en una fecha/hora con timezone local.

    Si ya tiene timezone, se conserva.
    """

    if value.tzinfo is None:
        return value.replace(
            tzinfo=LOCAL_TIMEZONE
        )

    return value

def as_utc(value: datetime) -> datetime:
    """Compara fechas de la base de datos sin importar si traen zona horaria."""

    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)

    return value.astimezone(timezone.utc)

@dataclass
class Schedule:
    start: int        # minutos desde medianoche
    end: int
    tolerance: int
    days: list[int]   # lunes = 0
    custom: bool


def default_schedule() -> Schedule:
    return Schedule(
        start=to_minutes(settings.work_start_time[:5]),
        end=to_minutes(getattr(settings, "work_end_time", "18:00")[:5]),
        tolerance=settings.late_grace_minutes,
        days=[0, 1, 2, 3, 4],
        custom=False,
    )


def schedule_from_profile(profile: "EmployeeProfile | None") -> Schedule:
    base = default_schedule()

    if not profile or profile.work_start_minutes is None:
        return base

    return Schedule(
        start=profile.work_start_minutes,
        end=profile.work_end_minutes or base.end,
        tolerance=profile.tolerance_minutes if profile.tolerance_minutes is not None else base.tolerance,
        days=[int(d) for d in (profile.work_days or "0,1,2,3,4").split(",") if d != ""],
        custom=True,
    )


def get_profile(db: Session, user_id: int) -> "EmployeeProfile | None":
    return db.scalar(select(EmployeeProfile).where(EmployeeProfile.user_id == user_id))


def get_schedule(db: Session, user_id: int) -> Schedule:
    return schedule_from_profile(get_profile(db, user_id))


def is_late_entry(recorded_at: datetime, schedule: Schedule | None = None) -> bool:
    schedule = schedule or default_schedule()
    local = normalize_datetime(recorded_at).astimezone(LOCAL_TIMEZONE)
    seconds = local.hour * 3600 + local.minute * 60 + local.second

    return seconds > (schedule.start + schedule.tolerance) * 60


# =========================================================
# AUTENTICACIÓN
# =========================================================


def create_access_token(user: User) -> str:
    expires_at = datetime.now(timezone.utc) + timedelta(
        days=settings.jwt_expiry_days
    )

    payload = {
        "sub": str(user.id),
        "email": user.email,
        "role": user.role,
        "exp": expires_at,
    }

    return jwt.encode(
        payload,
        settings.jwt_secret,
        algorithm=JWT_ALGORITHM,
    )


def get_user_from_token(
    token: str,
    db: Session,
) -> User:

    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Token inválido o expirado.",
        headers={
            "WWW-Authenticate": "Bearer"
        },
    )

    try:
        payload = jwt.decode(
            token,
            settings.jwt_secret,
            algorithms=[JWT_ALGORITHM],
        )

        user_id = payload.get("sub")

        if not user_id:
            raise credentials_exception

        user = db.get(
            User,
            int(user_id),
        )

        if not user:
            raise credentials_exception

        if not user.is_active:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="El usuario está desactivado.",
            )

        return user

    except (jwt.PyJWTError, ValueError):
        raise credentials_exception


def current_user(
    db: Session = Depends(SessionLocal),
) -> User:
    """
    Placeholder reemplazado por el header Authorization
    en la dependencia real.
    """

    raise HTTPException(
        status_code=500,
        detail="Configuración incorrecta de autenticación.",
    )


# =========================================================
# DEPENDENCIA DB + AUTH
# =========================================================


def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()


from fastapi.security import OAuth2PasswordBearer

oauth2_scheme = OAuth2PasswordBearer(
    tokenUrl="/auth/login"
)


def authenticated_user(
    token: str = Depends(oauth2_scheme),
    db: Session = Depends(get_db),
) -> User:

    return get_user_from_token(
        token,
        db,
    )


def admin_user(
    user: User = Depends(authenticated_user),
) -> User:

    if user.role != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Se requieren permisos de administrador.",
        )

    return user


# =========================================================
# ASISTENCIA
# =========================================================


def get_attendance_for_date(
    user_id: int,
    work_date: date,
    db: Session,
) -> dict:

    records = db.scalars(
        select(AttendanceRecord)
        .where(
            AttendanceRecord.user_id == user_id,
            AttendanceRecord.work_date == work_date,
        )
        .order_by(
            AttendanceRecord.recorded_at
        )
    ).all()

    result = {
        "entry": None,
        "exit": None,
    }

    for record in records:
        if record.kind == "entry":
            result["entry"] = record

        elif record.kind == "exit":
            result["exit"] = record

    return result


def get_today_attendance(
    user_id: int,
    db: Session,
) -> dict:

    return get_attendance_for_date(
        user_id,
        get_work_date(),
        db,
    )


def get_active_vacation(
    user_id: int,
    db: Session,
    work_date: date | None = None,
) -> VacationRequest | None:

    if work_date is None:
        work_date = get_work_date()

    return db.scalar(
        select(VacationRequest)
        .where(
            VacationRequest.user_id == user_id,
            VacationRequest.status == "approved",
            VacationRequest.start_date <= work_date,
            VacationRequest.end_date >= work_date,
        )
        .limit(1)
    )


def create_attendance_record(
    *,
    user: User,
    db: Session,
    kind: str,
    source: str,
    latitude: float | None = None,
    longitude: float | None = None,
    accuracy_meters: float | None = None,
    recorded_at: datetime | None = None,
    work_date: date | None = None,
    status_override: str | None = None,
    created_by_id: int | None = None,
    note: str | None = None,
) -> AttendanceRecord:

    

    # -----------------------------------------------------
    # 1. Validar tipo de asistencia
    # -----------------------------------------------------

    if kind not in {"entry", "exit"}:
        raise HTTPException(
            status_code=400,
            detail="El tipo de asistencia debe ser entry o exit.",
        )

    # -----------------------------------------------------
    # 2. Fecha/hora
    # -----------------------------------------------------

    if recorded_at is None:
        recorded_at = datetime.now(timezone.utc)

    recorded_at = normalize_datetime(
        recorded_at
    )

    if work_date is None:
        work_date = get_work_date(
            recorded_at
        )

    # -----------------------------------------------------
    # 3. Validar que el usuario esté activo
    # -----------------------------------------------------

    if not user.is_active:
        raise HTTPException(
            status_code=403,
            detail="El usuario está desactivado.",
        )

    # -----------------------------------------------------
    # 4. Revisar vacaciones
    # -----------------------------------------------------

    vacation = get_active_vacation(
        user.id,
        db,
        work_date,
    )

    if vacation:
        raise HTTPException(
            status_code=409,
            detail=(
                "El usuario tiene vacaciones aprobadas "
                "para esta fecha."
            ),
        )

    # -----------------------------------------------------
    # 5. Revisar asistencia existente
    # -----------------------------------------------------

    attendance = get_attendance_for_date(
        user.id,
        work_date,
        db,
    )

    if kind == "entry" and attendance["entry"]:
        raise HTTPException(
            status_code=409,
            detail="El usuario ya registró su entrada para este día.",
        )

    if kind == "exit" and attendance["exit"]:
        raise HTTPException(
            status_code=409,
            detail="El usuario ya registró su salida para este día.",
        )

    # -----------------------------------------------------
    # 6. No permitir salida sin entrada
    # -----------------------------------------------------

    if kind == "exit" and not attendance["entry"]:
        raise HTTPException(
            status_code=409,
            detail=(
                "No se puede registrar la salida "
                "porque no existe una entrada para este día."
            ),
        )

    # -----------------------------------------------------
    # 7. GPS para registros hechos por empleado
    # -----------------------------------------------------

    if source == "employee":

        if (
            latitude is None
            or longitude is None
            or accuracy_meters is None
        ):
            raise HTTPException(
                status_code=400,
                detail="La ubicación es obligatoria.",
            )

        distance = distance_meters(
            latitude,
            longitude,
            settings.company_latitude,
            settings.company_longitude,
        )

        if distance > settings.allowed_radius_meters:
            raise HTTPException(
                status_code=403,
                detail=(
                    "Estás fuera del área permitida "
                    "para registrar asistencia."
                ),
            )

    # -----------------------------------------------------
    # 8. Registro administrativo
    # -----------------------------------------------------

    else:

        latitude = (
            settings.company_latitude
            if latitude is None
            else latitude
        )

        longitude = (
            settings.company_longitude
            if longitude is None
            else longitude
        )

        accuracy_meters = (
            0
            if accuracy_meters is None
            else accuracy_meters
        )

        distance = 0.0

    # -----------------------------------------------------
    # 9. Determinar estado
    # -----------------------------------------------------

    if status_override is not None:

        if status_override not in {
            "normal",
            "late",
        }:
            raise HTTPException(
                status_code=400,
                detail=(
                    "El estado debe ser normal o late."
                ),
            )

        attendance_status = status_override

    else:

        if kind == "entry":
            attendance_status = (
                "late"
                if is_late_entry(recorded_at, get_schedule(db, user.id))
                else "normal"
            )
        else:
            attendance_status = "normal"

    # -----------------------------------------------------
    # 10. Crear registro
    # -----------------------------------------------------

    record = AttendanceRecord(
        user_id=user.id,
        kind=kind,
        work_date=work_date,
        recorded_at=recorded_at.astimezone(timezone.utc),
        latitude=latitude,
        longitude=longitude,
        accuracy_meters=accuracy_meters,
        distance_meters=distance,
        source=source,
        status=attendance_status,
        created_by_id=created_by_id,
        note=note,
    )

    db.add(record)

    try:
        db.commit()
        db.refresh(record)

    except Exception as exc:
        db.rollback()

        # Puede ser una violación de la restricción única
        # si dos peticiones llegaron prácticamente al mismo tiempo.
        raise HTTPException(
            status_code=409,
            detail=(
                "No se pudo registrar la asistencia. "
                "Es posible que ya exista un registro "
                "para ese tipo y fecha."
            ),
        ) from exc

    return record

def ensure_user_columns() -> None:
    existing = {c["name"] for c in inspect(engine).get_columns("users")}

    with engine.begin() as conn:
        if "hire_date" not in existing:
            conn.execute(text("ALTER TABLE users ADD COLUMN hire_date DATE"))

        if "avatar" not in existing:
            conn.execute(text("ALTER TABLE users ADD COLUMN avatar TEXT"))

def ensure_profile_columns() -> None:
    existing = {c["name"] for c in inspect(engine).get_columns("employee_profiles")}

    if "is_top_level" not in existing:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE employee_profiles ADD COLUMN is_top_level BOOLEAN DEFAULT FALSE"))

# =========================================================
# REPORTES
# =========================================================

REPORT_KEYS = [
    "attendance", "late", "absence", "vacation", "paid_leave",
    "unpaid_leave", "holiday", "rest", "incomplete",
]

DEFAULT_CODES = {
    "attendance": ("ASI", "Asistencia", False),
    "late": ("RET", "Retardo", True),
    "absence": ("FAL", "Falta injustificada", True),
    "vacation": ("VAC", "Vacaciones", True),
    "paid_leave": ("PCG", "Permiso con goce de sueldo", True),
    "unpaid_leave": ("SGS", "Permiso sin goce de sueldo", True),
    "holiday": ("FES", "Día festivo", False),
    "rest": ("DES", "Día de descanso", False),
    "incomplete": ("INC", "Registro incompleto (sin salida)", True),
}

CODE_COLORS = {
    "attendance": "E7F5EB",
    "late": "FFF1D6",
    "absence": "FDE2DE",
    "vacation": "DCEBFF",
    "paid_leave": "ECE3FF",
    "unpaid_leave": "FFE6D2",
    "holiday": "E3E8F0",
    "rest": "F1F3F6",
    "incomplete": "FCE4EF",
}

WEEKDAYS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

THIN = Side(style="thin", color="DFE6EF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def ensure_incidence_codes(db: Session) -> None:
    existing = set(db.scalars(select(IncidenceCode.key)).all())

    for key in REPORT_KEYS:
        if key not in existing:
            code, label, include = DEFAULT_CODES[key]
            db.add(IncidenceCode(key=key, code=code, label=label, include_in_incidences=include))


def load_codes(db: Session) -> list[IncidenceCodeOut]:
    rows = {c.key: c for c in db.scalars(select(IncidenceCode)).all()}
    result = []

    for key in REPORT_KEYS:
        code, label, include = DEFAULT_CODES[key]
        row = rows.get(key)

        result.append(
            IncidenceCodeOut(
                key=key,
                code=row.code if row else code,
                label=row.label if row else label,
                include_in_incidences=row.include_in_incidences if row else include,
                color=f"#{CODE_COLORS[key]}",
            )
        )

    return result


@dataclass
class ReportDay:
    user: User
    day: date
    key: str
    entry: datetime | None = None
    exit: datetime | None = None
    hours: float = 0.0
    late_minutes: int = 0
    incomplete: bool = False
    note: str | None = None


def build_report_days(db: Session, data: ReportRequest):
    """Clasifica cada día de cada empleado del periodo con una clave."""

    profiles = {p.user_id: p for p in db.scalars(select(EmployeeProfile)).all()}

    query = select(User).where(User.role == "employee").order_by(User.full_name)
    if data.user_id:
        query = query.where(User.id == data.user_id)

    users = [
        u
        for u in db.scalars(query).all()
        if not data.department
        or (profiles.get(u.id) and profiles[u.id].department == data.department)
    ]

    if not users:
        return [], profiles, []

    ids = [u.id for u in users]

    attendance: dict[tuple[int, date], dict[str, AttendanceRecord]] = {}

    for record in db.scalars(
        select(AttendanceRecord).where(
            AttendanceRecord.user_id.in_(ids),
            AttendanceRecord.work_date >= data.start,
            AttendanceRecord.work_date <= data.end,
        )
    ).all():
        attendance.setdefault((record.user_id, record.work_date), {})[record.kind] = record

    def approved(model):
        grouped: dict[int, list] = {}

        for item in db.scalars(
            select(model).where(
                model.user_id.in_(ids),
                model.status == "approved",
                model.start_date <= data.end,
                model.end_date >= data.start,
            )
        ).all():
            grouped.setdefault(item.user_id, []).append(item)

        return grouped

    vacations = approved(VacationRequest)
    permissions = approved(PermissionRequest)

    holidays = {
        h.holiday_date: h.name
        for h in db.scalars(
            select(Holiday).where(
                Holiday.holiday_date >= data.start,
                Holiday.holiday_date <= data.end,
            )
        ).all()
    }

    today = get_work_date()
    total_days = (data.end - data.start).days + 1

    included: list[User] = []
    result: list[ReportDay] = []

    for user in users:
        schedule = schedule_from_profile(profiles.get(user.id))
        user_days: list[ReportDay] = []
        worked = False

        for offset in range(total_days):
            day = data.start + timedelta(days=offset)

            if day > today or (user.hire_date and day < user.hire_date):
                continue

            records = attendance.get((user.id, day), {})
            entry = records.get("entry")
            exit_record = records.get("exit")

            if entry:
                worked = True
                late = entry.status == "late"
                local_entry = as_utc(entry.recorded_at).astimezone(LOCAL_TIMEZONE)

                hours = 0.0
                if exit_record:
                    hours = round(
                        (as_utc(exit_record.recorded_at) - as_utc(entry.recorded_at)).total_seconds() / 3600,
                        2,
                    )

                notes = [n for n in (entry.note, exit_record.note if exit_record else None) if n]

                user_days.append(
                    ReportDay(
                        user=user,
                        day=day,
                        key="late" if late else "attendance",
                        entry=entry.recorded_at,
                        exit=exit_record.recorded_at if exit_record else None,
                        hours=hours,
                        late_minutes=(
                            max(0, local_entry.hour * 60 + local_entry.minute - schedule.start) if late else 0
                        ),
                        incomplete=exit_record is None and day < today,
                        note="; ".join(notes) or None,
                    )
                )
                continue

            vacation = next((v for v in vacations.get(user.id, []) if v.start_date <= day <= v.end_date), None)
            permission = next((p for p in permissions.get(user.id, []) if p.start_date <= day <= p.end_date), None)

            if day in holidays:
                key, note = "holiday", holidays[day]
            elif day.weekday() not in schedule.days:
                key, note = "rest", None
            elif vacation:
                key, note = "vacation", vacation.reason
            elif permission:
                key = "paid_leave" if permission.kind == "paid" else "unpaid_leave"
                note = permission.reason
            elif day == today:
                continue  # el día en curso solo cuenta si ya hay registro
            else:
                key, note = "absence", None

            user_days.append(ReportDay(user=user, day=day, key=key, note=note))

        # Los inactivos solo aparecen si trabajaron en el periodo
        if user.is_active or worked:
            included.append(user)
            result.extend(user_days)

    return included, profiles, result


def summarize_report(users: list[User], days: list[ReportDay]) -> dict[int, dict]:
    stats = {
        u.id: {"counts": {k: 0 for k in REPORT_KEYS}, "hours": 0.0, "late": 0, "workdays": 0}
        for u in users
    }

    for item in days:
        s = stats[item.user.id]
        s["counts"][item.key] += 1

        if item.incomplete:
            s["counts"]["incomplete"] += 1

        s["hours"] += item.hours
        s["late"] += item.late_minutes

        if item.key not in ("rest", "holiday"):
            s["workdays"] += 1

    return stats


def employee_fields(user: User, profiles: dict) -> tuple[str, str, str]:
    profile = profiles.get(user.id)

    return (
        (profile.employee_code if profile else None) or "",
        (profile.department if profile else None) or "",
        (profile.position if profile else None) or "",
    )


def local_hhmm(value: datetime | None) -> str:
    return as_utc(value).astimezone(LOCAL_TIMEZONE).strftime("%H:%M") if value else ""


def style_header(cell, fill_hex: str) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill_hex)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def solid(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def write_table(sheet, headers: list[str], rows: list[list], header_hex: str, widths: list[int]) -> None:
    """Encabezado con color, bordes, filas alternadas, filtros y panel fijo."""

    for col, title in enumerate(headers, start=1):
        style_header(sheet.cell(row=1, column=col, value=title), header_hex)

    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            cell = sheet.cell(row=r, column=c, value=value)
            cell.border = BORDER

            if r % 2:
                cell.fill = solid("F7F9FC")

    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"

# =========================================================
# STARTUP
# =========================================================


@app.on_event("startup")
def startup():

    Base.metadata.create_all(
        bind=engine
    )

    ensure_user_columns()
    ensure_profile_columns()

    db = SessionLocal()

    try:

        admin = db.scalar(
            select(User)
            .where(
                User.email
                == settings.initial_admin_email
            )
        )

        if not admin:

            admin = User(
                email=settings.initial_admin_email,
                full_name="Administrador",
                password_hash=password_hash.hash(
                    settings.initial_admin_password
                ),
                role="admin",
                is_active=True,
            )

            db.add(admin)

        # -------------------------------------------------
        # Crear empleado inicial si no existe
        # -------------------------------------------------

        employees = [
           ("lili@integraprofesional.com","Lili"),
           ("cindy@integraprofesional.com","Cindy"),
           ("eleazar@integraprofesional.com","Eleazar"),
           ("angel@integraprofesional.com","Angel"),
           ("jesus@integraprofesional.com","Jesus"),
           ("norma@integraprofesional.com", "Norma"),
           ("raul@integraprofesional.com", "Raul")
        ]

        for email, full_name in employees:
            employee = db.scalar(
                select(User)
                .where(
                    User.email == email
                )
            )

            if not employee:

                employee = User(
                    email= email,
                    full_name= full_name,
                    password_hash=password_hash.hash(
                        settings.employee_initial_password
                    ),
                    role="employee",
                    is_active=True,
                )

                db.add(employee)


        get_branding(db)
        db.commit()

    finally:
        db.close()


# =========================================================
# HEALTH
# =========================================================


@app.get("/health")
def health():
    return {
        "status": "ok"
    }


# =========================================================
# AUTH
# =========================================================


@app.post(
    "/auth/login",
    response_model=TokenOut,
)
def login(
    data: LoginInput,
    db: Session = Depends(get_db),
):

    user = db.scalar(
        select(User)
        .where(
            User.email == str(data.email).lower()
        )
    )

    if not user:

        raise HTTPException(
            status_code=401,
            detail="Correo o contraseña incorrectos.",
        )

    if not password_hash.verify(
        data.password,
        user.password_hash,
    ):

        raise HTTPException(
            status_code=401,
            detail="Correo o contraseña incorrectos.",
        )

    if not user.is_active:

        raise HTTPException(
            status_code=403,
            detail="El usuario está desactivado.",
        )

    token = create_access_token(user)

    return TokenOut(
        access_token=token,
        token_type="bearer",
        user=user,
    )


@app.get(
    "/auth/me",
    response_model=UserOut,
)
def me(
    user: User = Depends(authenticated_user),
):
    return user


# =========================================================
# USERS
# =========================================================


@app.get(
    "/users",
    response_model=list[UserOut],
)
def list_users(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    return db.scalars(
        select(User)
        .order_by(User.full_name)
    ).all()


@app.post(
    "/users",
    response_model=UserOut,
)
def create_user(
    data: UserCreate,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    email = str(data.email).lower()

    existing = db.scalar(
        select(User)
        .where(User.email == email)
    )

    if existing:

        raise HTTPException(
            status_code=409,
            detail="Ya existe un usuario con ese correo.",
        )

    hire_date = data.hire_date 

    user = User(
        email=email,
        full_name=data.full_name,
        password_hash=password_hash.hash(
            settings.employee_initial_password
        ),
        role=data.role,
        is_active=True,
    )

    db.add(user)
    db.commit()
    db.refresh(user)

    return user


# =========================================================
# ATTENDANCE - EMPLOYEE
# =========================================================


@app.post(
    "/attendance/manual",
    response_model=AttendanceOut,
)
def manual_attendance(
    data: ManualAttendanceInput,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    employee = db.get(
        User,
        data.user_id,
    )

    if not employee:

        raise HTTPException(
            status_code=404,
            detail="Empleado no encontrado.",
        )

    if employee.role == "admin":

        raise HTTPException(
            status_code=400,
            detail="No se puede registrar asistencia manual para un administrador.",
        )

    today = get_work_date()

    if data.work_date > today:

        raise HTTPException(
            status_code=400,
            detail="No se puede registrar asistencia para una fecha futura.",
        )

    recorded_at = normalize_datetime(
        data.recorded_at
    )

    # Si el administrador manda una fecha diferente
    # a work_date, verificamos que realmente coincidan.
    recorded_local_date = get_work_date(
        recorded_at
    )

    if recorded_local_date != data.work_date:

        raise HTTPException(
            status_code=400,
            detail=(
                "La fecha de recorded_at no coincide "
                "con work_date."
            ),
        )

    record = create_attendance_record(
        user=employee,
        db=db,
        kind=data.kind,
        source="admin",
        recorded_at=recorded_at,
        work_date=data.work_date,
        status_override=data.status,
        created_by_id=admin.id,
        note=data.note,
    )

    create_notification(
        db,
        employee.id,
        "attendance_fix",
        "Se registró una asistencia a tu nombre",
        (
            f"Un administrador registró tu "
            f"{'entrada' if data.kind == 'entry' else 'salida'} del {data.work_date:%d/%m/%Y}."
        ),
        "attendance",
        f"manual-{record.id}",
    )
    db.commit()

    return record

@app.post(
    "/attendance/{kind}",
    response_model=AttendanceOut,
)
def register_attendance(
    kind: str,
    data: AttendanceInput,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    return create_attendance_record(
        user=user,
        db=db,
        kind=kind,
        source="employee",
        latitude=data.latitude,
        longitude=data.longitude,
        accuracy_meters=data.accuracy_meters,
    )


@app.get(
    "/attendance/mine",
    response_model=list[AttendanceOut],
)
def my_attendance(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    return db.scalars(
        select(AttendanceRecord)
        .where(
            AttendanceRecord.user_id == user.id
        )
        .order_by(
            AttendanceRecord.recorded_at.desc()
        )
    ).all()


@app.get(
    "/attendance/today",
    response_model=AttendanceTodayOut,
)
def attendance_today(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    today = get_work_date()

    attendance = get_attendance_for_date(
        user.id,
        today,
        db,
    )

    vacation = get_active_vacation(
        user.id,
        db,
        today,
    )

    return AttendanceTodayOut(
        work_date=today,
        entry=attendance["entry"],
        exit=attendance["exit"],
        is_on_vacation=vacation is not None,
    )


# =========================================================
# ATTENDANCE - ADMIN
# =========================================================


@app.get(
    "/attendance",
    response_model=list[AttendanceAdminOut],
)
def list_all_attendance(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    rows = db.execute(
        select(
            AttendanceRecord,
            User.full_name,
            User.email,
        )
        .join(
            User,
            AttendanceRecord.user_id == User.id,
        )
        .order_by(
            AttendanceRecord.work_date.desc(),
            AttendanceRecord.recorded_at.desc(),
        )
    ).all()

    result = []

    for record, user_name, user_email in rows:

        result.append(
            AttendanceAdminOut(
                id=record.id,
                user_id=record.user_id,
                kind=record.kind,
                work_date=record.work_date,
                recorded_at=record.recorded_at,
                distance_meters=record.distance_meters,
                accuracy_meters=record.accuracy_meters,
                source=record.source,
                status=record.status,
                note=record.note,
                user_name=user_name,
                user_email=user_email,
            )
        )

    return result


@app.get(
    "/attendance/today-overview",
    response_model=TodayOverviewOut,
)
def today_overview(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    now_local = get_local_now()
    today = now_local.date()
    now_minutes = now_local.hour * 60 + now_local.minute

    holiday = db.scalar(select(Holiday).where(Holiday.holiday_date == today))

    employees = db.scalars(
        select(User)
        .where(User.role == "employee", User.is_active.is_(True))
        .order_by(User.full_name)
    ).all()

    profiles = {p.user_id: p for p in db.scalars(select(EmployeeProfile)).all()}

    attendance: dict[int, dict[str, AttendanceRecord]] = {}
    for record in db.scalars(
        select(AttendanceRecord).where(AttendanceRecord.work_date == today)
    ).all():
        attendance.setdefault(record.user_id, {})[record.kind] = record

    def approved(model):
        return {
            item.user_id: item
            for item in db.scalars(
                select(model).where(
                    model.status == "approved",
                    model.start_date <= today,
                    model.end_date >= today,
                )
            ).all()
        }

    vacations = approved(VacationRequest)
    permissions = approved(PermissionRequest)

    registered, missing, on_vacation, on_permission = [], [], [], []
    pending = 0
    anyone_works = False

    for employee in employees:
        user_out = UserOut.model_validate(employee)
        schedule = schedule_from_profile(profiles.get(employee.id))
        works_today = holiday is None and today.weekday() in schedule.days
        anyone_works = anyone_works or works_today

        records = attendance.get(employee.id, {})
        entry = records.get("entry")
        exit_record = records.get("exit")

        if entry:
            registered.append(
                TodayEntryOut(
                    user=user_out,
                    entry_at=entry.recorded_at,
                    exit_at=exit_record.recorded_at if exit_record else None,
                    status=entry.status,
                )
            )
        elif employee.id in vacations:
            v = vacations[employee.id]
            on_vacation.append(TodayAbsenceOut(user=user_out, start_date=v.start_date, end_date=v.end_date))
        elif employee.id in permissions:
            p = permissions[employee.id]
            on_permission.append(
                TodayAbsenceOut(user=user_out, start_date=p.start_date, end_date=p.end_date, kind=p.kind)
            )
        elif works_today:
            if now_minutes > schedule.start + schedule.tolerance:
                missing.append(user_out)
            else:
                pending += 1

    registered.sort(key=lambda item: item.entry_at)

    return TodayOverviewOut(
        work_date=today,
        is_workday=anyone_works,
        holiday_name=holiday.name if holiday else None,
        pending_count=pending,
        registered=registered,
        missing=missing,
        on_vacation=on_vacation,
        on_permission=on_permission,
    )

@app.patch(
    "/attendance/{record_id}",
    response_model=AttendanceOut,
)
def edit_attendance(
    record_id: int,
    data: AttendanceEdit,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
    
):
    record = db.get(AttendanceRecord, record_id)

    

    if not record:
        raise HTTPException(status_code=404, detail="Registro no encontrado.")

    time_changed = False

    if data.recorded_at is not None:
        new_local = normalize_datetime(data.recorded_at)

        if get_work_date(new_local) != record.work_date:
            raise HTTPException(
                status_code=400,
                detail="La hora debe corresponder al mismo día del registro.",
            )

        other = db.scalar(
            select(AttendanceRecord).where(
                AttendanceRecord.user_id == record.user_id,
                AttendanceRecord.work_date == record.work_date,
                AttendanceRecord.kind == ("exit" if record.kind == "entry" else "entry"),
            )
        )

        new_utc = new_local.astimezone(timezone.utc)

        if other:
            if record.kind == "entry" and new_utc >= as_utc(other.recorded_at):
                raise HTTPException(status_code=400, detail="La entrada debe ser anterior a la salida.")

            if record.kind == "exit" and new_utc <= as_utc(other.recorded_at):
                raise HTTPException(status_code=400, detail="La salida debe ser posterior a la entrada.")

        record.recorded_at = new_utc
        time_changed = True

        if record.kind == "entry" and data.status is None:
            record.status = "late" if is_late_entry(new_local, get_schedule(db, record.user_id)) else "normal"

    if data.status is not None and record.kind == "entry":
        record.status = data.status

    if data.note is not None:
        record.note = data.note.strip() or None

    record.source = "admin"
    record.created_by_id = admin.id

    create_notification(
        db,
        record.user_id,
        "attendance_fix",
        "Se corrigió un registro de asistencia",
        (
            f"Un administrador actualizó tu "
            f"{'entrada' if record.kind == 'entry' else 'salida'} del {record.work_date:%d/%m/%Y}."
        ),
        "attendance",
        f"fix-{record.id}-{int(datetime.now(timezone.utc).timestamp())}",
    )

    db.commit()
    db.refresh(record)

    return record


@app.delete("/attendance/{record_id}")
def delete_attendance(
    record_id: int,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    record = db.get(AttendanceRecord, record_id)

    if not record:
        raise HTTPException(status_code=404, detail="Registro no encontrado.")

    if record.kind == "entry":
        has_exit = db.scalar(
            select(AttendanceRecord.id).where(
                AttendanceRecord.user_id == record.user_id,
                AttendanceRecord.work_date == record.work_date,
                AttendanceRecord.kind == "exit",
            )
        )

        if has_exit:
            raise HTTPException(
                status_code=409,
                detail="Elimina primero la salida de ese día.",
            )

    create_notification(
        db,
        record.user_id,
        "attendance_fix",
        "Se eliminó un registro de asistencia",
        (
            f"Un administrador eliminó tu "
            f"{'entrada' if record.kind == 'entry' else 'salida'} del {record.work_date:%d/%m/%Y}."
        ),
        "attendance",
        f"delete-{record.id}-{int(datetime.now(timezone.utc).timestamp())}",
    )

    db.delete(record)
    db.commit()

    return {"ok": True}


# =========================================================
# QR ATTENDANCE
# =========================================================


@app.post(
    "/qr-attendance",
    response_model=AttendanceOut,
)
def qr_attendance(
    data: QrAttendanceInput,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    if data.station != settings.attendance_station_code:

        raise HTTPException(
            status_code=403,
            detail="Estación QR no válida.",
        )

    attendance = get_today_attendance(
        user.id,
        db,
    )

    if not attendance["entry"]:
        kind = "entry"

    elif not attendance["exit"]:
        kind = "exit"

    else:

        raise HTTPException(
            status_code=409,
            detail=(
                "Ya registraste entrada y salida "
                "para este día."
            ),
        )

    return create_attendance_record(
        user=user,
        db=db,
        kind=kind,
        source="employee",
        latitude=data.latitude,
        longitude=data.longitude,
        accuracy_meters=data.accuracy_meters,
    )


# =========================================================
# QR SESSION
# =========================================================


@app.post(
    "/qr-sessions",
    response_model=QrOut,
)
def create_qr_session(
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    code = uuid4().hex

    expires_at = datetime.now(timezone.utc) + timedelta(
        minutes=10
    )

    session = QrSession(
        code=code,
        expires_at=expires_at,
        created_by_id=admin.id,
    )

    db.add(session)
    db.commit()
    db.refresh(session)

    return QrOut(
        code=session.code,
        expires_at=session.expires_at,
    )


# =========================================================
# VACACIONES - EMPLOYEE
# =========================================================


@app.post(
    "/vacations/requests",
    response_model=VacationRequestOut,
)
def create_vacation_request(
    data: VacationRequestCreate,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    today = get_work_date()

    if data.start_date < today:

        raise HTTPException(
            status_code=400,
            detail=(
                "La fecha inicial de vacaciones "
                "no puede estar en el pasado."
            ),
        )

    if data.end_date < data.start_date:

        raise HTTPException(
            status_code=400,
            detail=(
                "La fecha final no puede ser anterior "
                "a la fecha inicial."
            ),
        )

    overlapping = db.scalar(
        select(VacationRequest)
        .where(
            VacationRequest.user_id == user.id,
            VacationRequest.status.in_(
                ["pending", "approved"]
            ),
            VacationRequest.start_date
            <= data.end_date,
            VacationRequest.end_date
            >= data.start_date,
        )
        .limit(1)
    )

    if overlapping:

        raise HTTPException(
            status_code=409,
            detail=(
                "Ya existe una solicitud o vacaciones "
                "aprobadas que se cruzan con esas fechas."
            ),
        )

    if count_workdays(data.start_date, data.end_date) == 0:
        raise HTTPException(
            status_code=400,
            detail="El periodo no incluye días laborables.",
        )

    for year in range(data.start_date.year, data.end_date.year + 1):
        first = date(year, 1, 1)
        last = date(year, 12, 31)

        requested = count_workdays(
            max(data.start_date, first),
            min(data.end_date, last),
        )

        used = vacation_days_in_year(
            db, user.id, year, ["pending", "approved"]
        )

        available = max(
            vacation_entitlement(user.hire_date, year) - used,
            0,
        )

        if requested > available:
            raise HTTPException(
                status_code=409,
                detail=(
                    f"No tienes suficientes días en {year}: "
                    f"solicitas {requested} y tienes {available} disponibles."
                ),
            )

    vacation = VacationRequest(
        user_id=user.id,
        start_date=data.start_date,
        end_date=data.end_date,
        status="pending",
        request_type="employee",
        reason=data.reason,
    )

    db.add(vacation)
    db.commit()
    db.refresh(vacation)

    notify_admins(
        db, "vacation", "Nueva solicitud de vacaciones",
        f"{user.full_name} solicitó vacaciones del {vacation.start_date:%d/%m/%Y} al {vacation.end_date:%d/%m/%Y}.",
        "admin-requests", f"vacation-new-{vacation.id}",
    )
    db.commit()

    return vacation


@app.get(
    "/vacations/mine",
    response_model=list[VacationRequestOut],
)
def my_vacations(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    vacations = db.scalars(
        select(VacationRequest)
        .where(VacationRequest.user_id == user.id)
        .order_by(VacationRequest.start_date.desc())
    ).all()

    return [vacation_to_out(v, db) for v in vacations]


@app.post(
    "/vacations/{vacation_id}/cancel",
    response_model=VacationRequestOut,
)
def cancel_vacation(
    vacation_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    vacation = db.get(
        VacationRequest,
        vacation_id,
    )

    if not vacation:

        raise HTTPException(
            status_code=404,
            detail="Solicitud de vacaciones no encontrada.",
        )

    if vacation.user_id != user.id:

        raise HTTPException(
            status_code=403,
            detail="No puedes cancelar esta solicitud.",
        )

    if vacation.status != "pending":

        raise HTTPException(
            status_code=409,
            detail=(
                "Solo se pueden cancelar solicitudes pendientes."
            ),
        )

    vacation.status = "cancelled"
    vacation.updated_at = datetime.now(
        timezone.utc
    )

    db.commit()
    db.refresh(vacation)

    return vacation


def vacation_days_in_year(
    db: Session,
    user_id: int,
    year: int,
    statuses: list[str],
) -> int:

    first = date(year, 1, 1)
    last = date(year, 12, 31)

    rows = db.scalars(
        select(VacationRequest).where(
            VacationRequest.user_id == user_id,
            VacationRequest.status.in_(statuses),
            VacationRequest.start_date <= last,
            VacationRequest.end_date >= first,
        )
    ).all()

    return sum(
        count_workdays(max(r.start_date, first), min(r.end_date, last))
        for r in rows
    )


def vacation_entitlement(hire_date: date | None, year: int) -> int:
    if hire_date is None:
        return VACATION_BASE_DAYS

    return VACATION_BASE_DAYS + VACATION_YEARLY_INCREMENT * max(
        year - hire_date.year, 0
    )


def build_vacation_balance(
    db: Session,
    user: User,
    year: int,
) -> VacationBalanceOut:

    total = vacation_entitlement(user.hire_date, year)
    used = vacation_days_in_year(db, user.id, year, ["approved"])
    pending = vacation_days_in_year(db, user.id, year, ["pending"])

    return VacationBalanceOut(
        year=year,
        total_days=total,
        used_days=used,
        pending_days=pending,
        available_days=max(total - used - pending, 0),
    )

def vacation_to_out(
    vacation: VacationRequest,
    db: Session,
) -> VacationRequestOut:

    out = VacationRequestOut.model_validate(vacation)

    if vacation.reviewed_by_id:
        reviewer = db.get(User, vacation.reviewed_by_id)
        out.reviewed_by_name = reviewer.full_name if reviewer else None

    return out


@app.get(
    "/vacations/balance",
    response_model=VacationBalanceOut,
)
def my_vacation_balance(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    return build_vacation_balance(db, user, get_work_date().year)

# =========================================================
# VACACIONES - ADMIN
# =========================================================


@app.get(
    "/vacations/requests",
    response_model=list[VacationRequestAdminOut],
)
def list_vacation_requests(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    rows = db.execute(
        select(
            VacationRequest,
            User.full_name,
            User.email,
        )
        .join(
            User,
            VacationRequest.user_id == User.id,
        )
        .order_by(
            VacationRequest.created_at.desc()
        )
    ).all()

    result = []

    for vacation, user_name, user_email in rows:

        result.append(
            VacationRequestAdminOut(
                id=vacation.id,
                user_id=vacation.user_id,
                start_date=vacation.start_date,
                end_date=vacation.end_date,
                status=vacation.status,
                request_type=vacation.request_type,
                reason=vacation.reason,
                admin_note=vacation.admin_note,
                reviewed_by_id=vacation.reviewed_by_id,
                reviewed_at=vacation.reviewed_at,
                created_at=vacation.created_at,
                updated_at=vacation.updated_at,
                user_name=user_name,
                user_email=user_email,
            )
        )

    return result


@app.post(
    "/vacations/{vacation_id}/approve",
    response_model=VacationRequestOut,
)
def approve_vacation(
    vacation_id: int,
    data: VacationReviewInput,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    vacation = db.get(
        VacationRequest,
        vacation_id,
    )

    if not vacation:

        raise HTTPException(
            status_code=404,
            detail="Solicitud de vacaciones no encontrada.",
        )

    if vacation.status != "pending":

        raise HTTPException(
            status_code=409,
            detail=(
                "Solo se pueden aprobar solicitudes pendientes."
            ),
        )

    overlapping = db.scalar(
        select(VacationRequest)
        .where(
            VacationRequest.user_id == vacation.user_id,
            VacationRequest.id != vacation.id,
            VacationRequest.status == "approved",
            VacationRequest.start_date
            <= vacation.end_date,
            VacationRequest.end_date
            >= vacation.start_date,
        )
        .limit(1)
    )

    if overlapping:

        raise HTTPException(
            status_code=409,
            detail=(
                "El empleado ya tiene vacaciones aprobadas "
                "que se cruzan con estas fechas."
            ),
        )

    vacation.status = "approved"
    vacation.admin_note = data.admin_note
    vacation.reviewed_by_id = admin.id
    vacation.reviewed_at = datetime.now(
        timezone.utc
    )

    create_notification(
        db, vacation.user_id, "vacation", "Vacaciones aprobadas",
        f"Tu solicitud del {vacation.start_date:%d/%m/%Y} al {vacation.end_date:%d/%m/%Y} fue aprobada.",
        "vacations", f"vacation-{vacation.id}-approved",
    )

    db.commit()
    db.refresh(vacation)

    return vacation


@app.post(
    "/vacations/{vacation_id}/reject",
    response_model=VacationRequestOut,
)
def reject_vacation(
    vacation_id: int,
    data: VacationReviewInput,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    vacation = db.get(
        VacationRequest,
        vacation_id,
    )

    if not vacation:

        raise HTTPException(
            status_code=404,
            detail="Solicitud de vacaciones no encontrada.",
        )

    if vacation.status != "pending":

        raise HTTPException(
            status_code=409,
            detail=(
                "Solo se pueden rechazar solicitudes pendientes."
            ),
        )

    vacation.status = "rejected"
    vacation.admin_note = data.admin_note
    vacation.reviewed_by_id = admin.id
    vacation.reviewed_at = datetime.now(
        timezone.utc
    )

    create_notification(
        db, vacation.user_id, "vacation", "Vacaciones rechazadas",
        f"Tu solicitud del {vacation.start_date:%d/%m/%Y} al {vacation.end_date:%d/%m/%Y} fue rechazada.",
        "vacations", f"vacation-{vacation.id}-rejected",
    )

    db.commit()
    db.refresh(vacation)

    return vacation


@app.post(
    "/vacations/admin",
    response_model=VacationRequestOut,
)
def create_admin_vacation(
    data: AdminVacationCreate,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    employee = db.get(
        User,
        data.user_id,
    )

    if not employee:

        raise HTTPException(
            status_code=404,
            detail="Empleado no encontrado.",
        )

    if employee.role == "admin":

        raise HTTPException(
            status_code=400,
            detail=(
                "No se pueden registrar vacaciones "
                "para otro administrador."
            ),
        )

    today = get_work_date()

    if data.start_date < today:

        raise HTTPException(
            status_code=400,
            detail=(
                "La fecha inicial de vacaciones "
                "no puede estar en el pasado."
            ),
        )

    overlapping = db.scalar(
        select(VacationRequest)
        .where(
            VacationRequest.user_id == data.user_id,
            VacationRequest.status.in_(
                ["pending", "approved"]
            ),
            VacationRequest.start_date
            <= data.end_date,
            VacationRequest.end_date
            >= data.start_date,
        )
        .limit(1)
    )

    if overlapping:

        raise HTTPException(
            status_code=409,
            detail=(
                "Ya existe una solicitud o vacaciones "
                "que se cruzan con esas fechas."
            ),
        )

    vacation = VacationRequest(
        user_id=data.user_id,
        start_date=data.start_date,
        end_date=data.end_date,
        status="approved",
        request_type="admin",
        reason=data.reason,
        admin_note=data.admin_note,
        reviewed_by_id=admin.id,
        reviewed_at=datetime.now(timezone.utc),
    )

    db.add(vacation)
    db.commit()
    db.refresh(vacation)

    return vacation

# =========================================================
# PERMISOS
# =========================================================


def permission_to_out(
    permission: PermissionRequest,
    db: Session,
) -> PermissionRequestOut:

    out = PermissionRequestOut.model_validate(permission)

    if permission.reviewed_by_id:
        reviewer = db.get(User, permission.reviewed_by_id)
        out.reviewed_by_name = reviewer.full_name if reviewer else None

    return out


@app.post(
    "/permissions",
    response_model=PermissionRequestOut,
)
def create_permission_request(
    data: PermissionRequestCreate,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    if data.start_date < get_work_date():
        raise HTTPException(
            status_code=400,
            detail="La fecha inicial no puede estar en el pasado.",
        )

    overlapping_permission = db.scalar(
        select(PermissionRequest)
        .where(
            PermissionRequest.user_id == user.id,
            PermissionRequest.status.in_(["pending", "approved"]),
            PermissionRequest.start_date <= data.end_date,
            PermissionRequest.end_date >= data.start_date,
        )
        .limit(1)
    )

    if overlapping_permission:
        raise HTTPException(
            status_code=409,
            detail=(
                "Ya existe un permiso pendiente o aprobado "
                "que se cruza con esas fechas."
            ),
        )

    overlapping_vacation = db.scalar(
        select(VacationRequest)
        .where(
            VacationRequest.user_id == user.id,
            VacationRequest.status == "approved",
            VacationRequest.start_date <= data.end_date,
            VacationRequest.end_date >= data.start_date,
        )
        .limit(1)
    )

    if overlapping_vacation:
        raise HTTPException(
            status_code=409,
            detail="Tienes vacaciones aprobadas en esas fechas.",
        )

    permission = PermissionRequest(
        user_id=user.id,
        kind=data.kind,
        start_date=data.start_date,
        end_date=data.end_date,
        reason=data.reason,
        status="pending",
    )

    db.add(permission)
    db.commit()
    db.refresh(permission)

    notify_admins(
        db, "permission", "Nueva solicitud de permiso",
        f"{user.full_name} solicitó un permiso del {permission.start_date:%d/%m/%Y} al {permission.end_date:%d/%m/%Y}.",
        "admin-requests", f"permission-new-{permission.id}",
    )
    db.commit()

    return permission_to_out(permission, db)


@app.get(
    "/permissions/mine",
    response_model=list[PermissionRequestOut],
)
def my_permissions(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    permissions = db.scalars(
        select(PermissionRequest)
        .where(PermissionRequest.user_id == user.id)
        .order_by(PermissionRequest.created_at.desc())
    ).all()

    return [permission_to_out(p, db) for p in permissions]


@app.post(
    "/permissions/{permission_id}/cancel",
    response_model=PermissionRequestOut,
)
def cancel_permission(
    permission_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    permission = db.get(PermissionRequest, permission_id)

    if not permission:
        raise HTTPException(
            status_code=404,
            detail="Solicitud de permiso no encontrada.",
        )

    if permission.user_id != user.id:
        raise HTTPException(
            status_code=403,
            detail="No puedes cancelar esta solicitud.",
        )

    if permission.status != "pending":
        raise HTTPException(
            status_code=409,
            detail="Solo se pueden cancelar solicitudes pendientes.",
        )

    permission.status = "cancelled"
    db.commit()
    db.refresh(permission)

    return permission_to_out(permission, db)


@app.get(
    "/permissions/requests",
    response_model=list[PermissionRequestAdminOut],
)
def list_permission_requests(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    rows = db.execute(
        select(
            PermissionRequest,
            User.full_name,
            User.email,
        )
        .join(User, PermissionRequest.user_id == User.id)
        .order_by(PermissionRequest.created_at.desc())
    ).all()

    result = []

    for permission, user_name, user_email in rows:
        base = permission_to_out(permission, db)

        result.append(
            PermissionRequestAdminOut(
                **base.model_dump(),
                user_name=user_name,
                user_email=user_email,
            )
        )

    return result


def review_permission(
    permission_id: int,
    data: PermissionReviewInput,
    admin: User,
    db: Session,
    new_status: str,
) -> PermissionRequestOut:

    permission = db.get(PermissionRequest, permission_id)

    if not permission:
        raise HTTPException(
            status_code=404,
            detail="Solicitud de permiso no encontrada.",
        )

    if permission.status != "pending":
        raise HTTPException(
            status_code=409,
            detail="Solo se pueden resolver solicitudes pendientes.",
        )

    permission.status = new_status
    permission.admin_note = data.admin_note
    permission.reviewed_by_id = admin.id
    permission.reviewed_at = datetime.now(timezone.utc)

    create_notification(
        db, permission.user_id, "permission",
        "Permiso aprobado" if new_status == "approved" else "Permiso rechazado",
        f"Tu permiso del {permission.start_date:%d/%m/%Y} al {permission.end_date:%d/%m/%Y} fue "
        f"{'aprobado' if new_status == 'approved' else 'rechazado'}.",
        "attendance", f"permission-{permission.id}-{new_status}",
    )

    db.commit()
    db.refresh(permission)

    return permission_to_out(permission, db)


@app.post(
    "/permissions/{permission_id}/approve",
    response_model=PermissionRequestOut,
)
def approve_permission(
    permission_id: int,
    data: PermissionReviewInput,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    return review_permission(permission_id, data, admin, db, "approved")


@app.post(
    "/permissions/{permission_id}/reject",
    response_model=PermissionRequestOut,
)
def reject_permission(
    permission_id: int,
    data: PermissionReviewInput,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    return review_permission(permission_id, data, admin, db, "rejected")


# =========================================================
# PERFIL, NOTIFICACIONES Y SEGURIDAD
# =========================================================


def get_or_create_preferences(
    db: Session,
    user_id: int,
) -> UserPreferences:

    prefs = db.scalar(
        select(UserPreferences).where(
            UserPreferences.user_id == user_id
        )
    )

    if not prefs:
        prefs = UserPreferences(user_id=user_id)
        db.add(prefs)
        db.commit()
        db.refresh(prefs)

    return prefs


def profile_out(user: User, prefs: UserPreferences) -> ProfileOut:
    return ProfileOut(
        full_name=user.full_name,
        email=user.email,
        phone=prefs.phone,
    )


@app.get(
    "/profile",
    response_model=ProfileOut,
)
def read_my_profile(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    return profile_out(user, get_or_create_preferences(db, user.id))


@app.put(
    "/profile",
    response_model=ProfileOut,
)
def update_profile(
    data: ProfileUpdate,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    prefs = get_or_create_preferences(db, user.id)

    user.full_name = data.full_name.strip()

    phone = (data.phone or "").strip()
    prefs.phone = phone or None

    db.commit()
    db.refresh(user)
    db.refresh(prefs)

    return profile_out(user, prefs)


@app.get(
    "/profile/notifications",
    response_model=NotificationSettings,
)
def get_notifications(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    return get_or_create_preferences(db, user.id)


@app.put(
    "/profile/notifications",
    response_model=NotificationSettings,
)
def update_notifications(
    data: NotificationSettings,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    prefs = get_or_create_preferences(db, user.id)

    prefs.notify_attendance = data.notify_attendance
    prefs.notify_vacations = data.notify_vacations
    prefs.notify_permissions = data.notify_permissions
    prefs.notify_company = data.notify_company
    prefs.notify_weekly = data.notify_weekly

    db.commit()
    db.refresh(prefs)

    return prefs


@app.post("/auth/change-password")
def change_password(
    data: PasswordChange,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    if not password_hash.verify(
        data.current_password,
        user.password_hash,
    ):
        # 400 y no 401, para no confundir con sesión expirada
        raise HTTPException(
            status_code=400,
            detail="La contraseña actual es incorrecta.",
        )

    if data.new_password == data.current_password:
        raise HTTPException(
            status_code=400,
            detail="La nueva contraseña debe ser diferente a la actual.",
        )

    user.password_hash = password_hash.hash(data.new_password)
    db.commit()

    return {"ok": True}


def generate_temporary_password() -> str:
    return secrets.token_urlsafe(9)  # 12 caracteres


@app.post(
    "/users",
    response_model=UserCreatedOut,
)
def create_user(
    data: UserCreate,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    email = str(data.email).lower()

    existing = db.scalar(
        select(User).where(User.email == email)
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Ya existe un usuario con ese correo.",
        )

    temporary_password = generate_temporary_password()

    user = User(
        email=email,
        full_name=data.full_name.strip(),
        password_hash=password_hash.hash(temporary_password),
        role=data.role,
        is_active=True,
    )

    db.add(user)
    db.commit()
    db.refresh(user)

    return UserCreatedOut(
        **UserOut.model_validate(user).model_dump(),
        temporary_password=temporary_password,
    )


@app.patch(
    "/users/{user_id}",
    response_model=UserOut,
)
def update_user(
    user_id: int,
    data: UserUpdate,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):

    target = db.get(User, user_id)

    if not target:
        raise HTTPException(status_code=404, detail="Usuario no encontrado.")

    if target.id == admin.id:
        if data.is_active is False or (data.role is not None and data.role != "admin"):
            raise HTTPException(
                status_code=400,
                detail="No puedes desactivarte ni quitarte el rol de administrador a ti mismo.",
            )

        if data.password is not None:
            raise HTTPException(
                status_code=400,
                detail="Para cambiar tu propia contraseña usa Configuración.",
            )

    if data.full_name is not None:
        target.full_name = data.full_name.strip()

    if data.role is not None:
        target.role = data.role

    if data.is_active is not None:
        target.is_active = data.is_active

    if data.hire_date is not None:
        target.hire_date = data.hire_date

    if data.password is not None:
        target.password_hash = password_hash.hash(data.password)

    db.commit()
    db.refresh(target)

    return target
# AVATAR

def validate_avatar(value: str, max_bytes: int = MAX_AVATAR_BYTES) -> None:
    for prefix, magic in AVATAR_FORMATS.items():
        if not value.startswith(prefix):
            continue

        try:
            raw = base64.b64decode(value[len(prefix):], validate=True)
        except Exception:
            break

        if len(raw) <= max_bytes and raw.startswith(magic):
            return

        break

    raise HTTPException(
        status_code=400,
        detail=f"La imagen no es válida. Usa JPG, PNG o WebP de máximo {max_bytes // 1000} KB.",
    )

@app.put(
    "/profile/avatar",
    response_model=UserOut,
)
def set_avatar(
    data: AvatarInput,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    validate_avatar(data.avatar)

    user.avatar = data.avatar
    db.commit()
    db.refresh(user)

    return user


@app.delete(
    "/profile/avatar",
    response_model=UserOut,
)
def remove_avatar(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    user.avatar = None
    db.commit()
    db.refresh(user)

    return user


def get_branding(db: Session) -> CompanySettings:
    row = db.get(CompanySettings, 1)

    if not row:
        row = CompanySettings(id=1)
        db.add(row)
        db.commit()
        db.refresh(row)

    return row


@app.get("/branding", response_model=BrandingOut)
def read_branding(db: Session = Depends(get_db)):
    # Público: el login necesita logo y colores antes de iniciar sesión
    return get_branding(db)


@app.put("/branding", response_model=BrandingOut)
def update_branding(
    data: BrandingUpdate,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    row = get_branding(db)

    row.company_name = data.company_name.strip()
    row.primary_color = data.primary_color.lower()
    row.sidebar_color = data.sidebar_color.lower()

    db.commit()
    db.refresh(row)

    return row


@app.put("/branding/logo", response_model=BrandingOut)
def set_logo(
    data: AvatarInput,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    validate_avatar(data.avatar)  # mismas reglas que la foto de perfil

    row = get_branding(db)
    row.logo = data.avatar
    db.commit()
    db.refresh(row)

    return row


@app.delete("/branding/logo", response_model=BrandingOut)
def remove_logo(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    row = get_branding(db)
    row.logo = None
    db.commit()
    db.refresh(row)

    return row

# =========================================================
# NOTIFICACIONES
# =========================================================

PREF_BY_KIND = {
    "missing_exit": "notify_attendance",
    "vacation": "notify_vacations",
    "permission": "notify_permissions",
}


def create_notification(
    db: Session,
    user_id: int,
    kind: str,
    title: str,
    body: str,
    page: str,
    dedupe_key: str,
) -> None:
    """Agrega la notificación a la sesión (el commit lo hace quien llama)."""

    pref = PREF_BY_KIND.get(kind)

    if pref and not getattr(get_or_create_preferences(db, user_id), pref):
        return

    exists = db.scalar(
        select(Notification.id).where(
            Notification.user_id == user_id,
            Notification.dedupe_key == dedupe_key,
        )
    )

    if not exists:
        db.add(
            Notification(
                user_id=user_id,
                kind=kind,
                title=title,
                body=body,
                page=page,
                dedupe_key=dedupe_key,
            )
        )


def notify_admins(
    db: Session,
    kind: str,
    title: str,
    body: str,
    page: str,
    dedupe_key: str,
) -> None:
    admin_ids = db.scalars(
        select(User.id).where(
            User.role == "admin",
            User.is_active.is_(True),
        )
    ).all()

    for admin_id in admin_ids:
        create_notification(db, admin_id, kind, title, body, page, dedupe_key)


def sync_attendance_notifications(
    db: Session,
    user: User,
    prefs: UserPreferences,
) -> None:
    """Avisa de entradas sin salida y retira el aviso si ya se corrigió."""

    if not prefs.notify_attendance:
        return

    now_local = get_local_now()
    today = now_local.date()

    hours, minutes = map(
        int,
        getattr(settings, "work_end_time", "18:00").split(":")[:2],
    )

    schedule = get_schedule(db, user.id)
    today_cutoff = now_local.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(
        minutes=schedule.end + 30
    )

    by_date: dict[date, dict[str, AttendanceRecord]] = {}

    for record in db.scalars(
        select(AttendanceRecord).where(
            AttendanceRecord.user_id == user.id,
            AttendanceRecord.work_date >= today - timedelta(days=14),
        )
    ).all():
        by_date.setdefault(record.work_date, {})[record.kind] = record

    existing = {
        n.dedupe_key: n
        for n in db.scalars(
            select(Notification).where(
                Notification.user_id == user.id,
                Notification.kind == "missing_exit",
            )
        ).all()
    }

    for work_date, kinds in by_date.items():
        key = f"missing-exit-{work_date.isoformat()}"

        missing = (
            "entry" in kinds
            and "exit" not in kinds
            and (work_date < today or now_local > today_cutoff)
        )

        if missing and key not in existing:
            db.add(
                Notification(
                    user_id=user.id,
                    kind="missing_exit",
                    title="No registraste tu salida",
                    body=(
                        f"El {work_date:%d/%m/%Y} registraste tu entrada pero no tu salida. "
                        "Pide a un administrador que la corrija."
                        if work_date < today
                        else "Ya terminó tu jornada y no has registrado tu salida."
                    ),
                    page="attendance",
                    dedupe_key=key,
                )
            )

        elif not missing and key in existing:
            db.delete(existing[key])

    db.commit()


def sync_weekly_summary(
    db: Session,
    user: User,
    prefs: UserPreferences,
) -> None:
    """Crea el resumen de la semana anterior (lunes a domingo), una sola vez."""

    if not prefs.notify_weekly or user.role != "employee":
        return

    today = get_work_date()
    monday = today - timedelta(days=today.weekday() + 7)
    sunday = monday + timedelta(days=6)
    key = f"weekly-{monday.isoformat()}"
    schedule = get_schedule(db, user.id)

    already = db.scalar(
        select(Notification.id).where(
            Notification.user_id == user.id,
            Notification.dedupe_key == key,
        )
    )

    if already or (user.hire_date and user.hire_date > sunday):
        return

    by_date: dict[date, dict[str, AttendanceRecord]] = {}

    for record in db.scalars(
        select(AttendanceRecord).where(
            AttendanceRecord.user_id == user.id,
            AttendanceRecord.work_date >= monday,
            AttendanceRecord.work_date <= sunday,
        )
    ).all():
        by_date.setdefault(record.work_date, {})[record.kind] = record

    holidays = set(
        db.scalars(
            select(Holiday.holiday_date).where(
                Holiday.holiday_date >= monday,
                Holiday.holiday_date <= sunday,
            )
        ).all()
    )

    away = [
        (item.start_date, item.end_date)
        for model in (VacationRequest, PermissionRequest)
        for item in db.scalars(
            select(model).where(
                model.user_id == user.id,
                model.status == "approved",
                model.start_date <= sunday,
                model.end_date >= monday,
            )
        ).all()
    ]

    punctual = late = absent = 0
    seconds = 0.0

    for offset in range(7):
        day = monday + timedelta(days=offset)
        records = by_date.get(day)

        if records and "entry" in records:
            if records["entry"].status == "late":
                late += 1
            else:
                punctual += 1

            if "exit" in records:
                seconds += (
                    as_utc(records["exit"].recorded_at)
                    - as_utc(records["entry"].recorded_at)
                ).total_seconds()

        elif (
            day.weekday() in schedule.days
            and day not in holidays
            and not any(start <= day <= end for start, end in away)
            and (not user.hire_date or day >= user.hire_date)
        ):
            absent += 1

    if punctual + late + absent == 0:
        return

    create_notification(
        db,
        user.id,
        "weekly",
        "Resumen de tu semana",
        (
            f"Del {monday:%d/%m} al {sunday:%d/%m}: {punctual} puntuales, "
            f"{late} con retardo, {absent} faltas y {seconds / 3600:.1f} h trabajadas."
        ),
        "history",
        key,
    )

@app.get(
    "/notifications",
    response_model=NotificationsOut,
)
def list_notifications(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):

    prefs = get_or_create_preferences(db, user.id)

    try:
        sync_attendance_notifications(db, user, prefs)
        sync_weekly_summary(db, user, prefs)
        sync_meeting_reminders(db, user)
        db.commit()
    except Exception:
        db.rollback()

    db.execute(
        delete(Notification).where(
            Notification.user_id == user.id,
            Notification.created_at < datetime.now(timezone.utc) - timedelta(days=60),
        )
    )
    db.commit()

    items = db.scalars(
        select(Notification)
        .where(Notification.user_id == user.id)
        .order_by(Notification.created_at.desc(), Notification.id.desc())
        .limit(30)
    ).all()

    unread = db.scalar(
        select(func.count())
        .select_from(Notification)
        .where(
            Notification.user_id == user.id,
            Notification.read_at.is_(None),
        )
    )

    return NotificationsOut(
        items=[NotificationOut.model_validate(n) for n in items],
        unread=unread or 0,
    )


@app.post("/notifications/read-all")
def read_all_notifications(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    db.execute(
        update(Notification)
        .where(
            Notification.user_id == user.id,
            Notification.read_at.is_(None),
        )
        .values(read_at=datetime.now(timezone.utc))
    )
    db.commit()

    return {"ok": True}


@app.post("/notifications/{notification_id}/read")
def read_notification(
    notification_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    db.execute(
        update(Notification)
        .where(
            Notification.id == notification_id,
            Notification.user_id == user.id,
            Notification.read_at.is_(None),
        )
        .values(read_at=datetime.now(timezone.utc))
    )
    db.commit()

    return {"ok": True}

# =========================================================
# ESPACIOS Y RESERVAS
# =========================================================


def space_to_out(space: Space) -> SpaceOut:
    return SpaceOut(
        id=space.id,
        name=space.name,
        kind=space.kind,
        capacity=space.capacity,
        location=space.location,
        description=space.description,
        equipment=[item for item in space.equipment.split("|") if item],
        open_time=to_hhmm(space.open_minutes),
        close_time=to_hhmm(space.close_minutes),
        open_days=[int(d) for d in space.open_days.split(",") if d != ""],
        is_active=space.is_active,
    )


def apply_space(space: Space, data: SpaceIn) -> None:
    space.name = data.name.strip()
    space.kind = data.kind.strip()
    space.capacity = data.capacity
    space.location = (data.location or "").strip() or None
    space.description = (data.description or "").strip() or None
    space.equipment = "|".join(data.equipment)
    space.open_minutes = to_minutes(data.open_time)
    space.close_minutes = to_minutes(data.close_time)
    space.open_days = ",".join(str(day) for day in data.open_days)
    space.is_active = data.is_active


def booking_to_out(
    booking: SpaceBooking,
    space_name: str,
    user_name: str,
) -> BookingOut:
    return BookingOut(
        id=booking.id,
        space_id=booking.space_id,
        space_name=space_name,
        user_id=booking.user_id,
        user_name=user_name,
        title=booking.title,
        notes=booking.notes,
        booking_date=booking.booking_date,
        start_time=to_hhmm(booking.start_minutes),
        end_time=to_hhmm(booking.end_minutes),
        status=booking.status,
    )

def space_conflict(
    db: Session,
    space_id: int,
    day: date,
    start: int,
    end: int,
    exclude_booking_id: int | None = None,
) -> SpaceBooking | None:
    query = select(SpaceBooking).where(
        SpaceBooking.space_id == space_id,
        SpaceBooking.booking_date == day,
        SpaceBooking.status == "confirmed",
        SpaceBooking.start_minutes < end,
        SpaceBooking.end_minutes > start,
    )

    if exclude_booking_id:
        query = query.where(SpaceBooking.id != exclude_booking_id)

    return db.scalar(query.limit(1))

def book_space(
    db: Session,
    user: User,
    space_id: int,
    day: date,
    start: int,
    end: int,
    title: str,
    notes: str | None = None,
) -> tuple[SpaceBooking, Space]:
    """Valida y agrega una reserva a la sesión (el commit lo hace quien llama)."""

    space = db.get(Space, space_id)

    if not space or not space.is_active:
        raise HTTPException(status_code=404, detail="Este espacio no está disponible.")

    if end <= start:
        raise HTTPException(status_code=400, detail="La hora de fin debe ser posterior a la de inicio.")

    if start % 15 or end % 15:
        raise HTTPException(status_code=400, detail="Las reservas son en bloques de 15 minutos.")

    if start < space.open_minutes or end > space.close_minutes:
        raise HTTPException(
            status_code=400,
            detail=f"{space.name} está disponible de {to_hhmm(space.open_minutes)} a {to_hhmm(space.close_minutes)}.",
        )

    if day.weekday() not in [int(d) for d in space.open_days.split(",") if d != ""]:
        raise HTTPException(status_code=400, detail=f"{space.name} no abre ese día.")

    now_local = get_local_now()

    if day < now_local.date() or (day == now_local.date() and start < now_local.hour * 60 + now_local.minute):
        raise HTTPException(status_code=400, detail="No puedes reservar en el pasado.")

    conflict = space_conflict(db, space.id, day, start, end)

    if conflict:
        raise HTTPException(
            status_code=409,
            detail=(
                f"{space.name} ya está reservada de {to_hhmm(conflict.start_minutes)} "
                f"a {to_hhmm(conflict.end_minutes)}."
            ),
        )

    booking = SpaceBooking(
        space_id=space.id,
        user_id=user.id,
        title=title,
        notes=notes,
        booking_date=day,
        start_minutes=start,
        end_minutes=end,
    )

    db.add(booking)
    db.flush()

    return booking, space


@app.get("/spaces", response_model=list[SpaceOut])
def list_spaces(
    include_inactive: bool = False,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    query = select(Space).order_by(Space.name)

    if not (include_inactive and user.role == "admin"):
        query = query.where(Space.is_active.is_(True))

    return [space_to_out(space) for space in db.scalars(query).all()]


@app.post("/spaces", response_model=SpaceOut)
def create_space(
    data: SpaceIn,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    space = Space()
    apply_space(space, data)

    db.add(space)
    db.commit()
    db.refresh(space)

    return space_to_out(space)


@app.put("/spaces/{space_id}", response_model=SpaceOut)
def update_space(
    space_id: int,
    data: SpaceIn,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    space = db.get(Space, space_id)

    if not space:
        raise HTTPException(status_code=404, detail="Espacio no encontrado.")

    apply_space(space, data)
    db.commit()
    db.refresh(space)

    return space_to_out(space)


@app.delete("/spaces/{space_id}")
def delete_space(
    space_id: int,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    space = db.get(Space, space_id)

    if not space:
        raise HTTPException(status_code=404, detail="Espacio no encontrado.")

    has_bookings = db.scalar(
        select(SpaceBooking.id).where(SpaceBooking.space_id == space_id).limit(1)
    )

    if has_bookings:
        raise HTTPException(
            status_code=409,
            detail="Este espacio ya tiene reservas. Desactívalo en lugar de eliminarlo.",
        )

    db.delete(space)
    db.commit()

    return {"ok": True}


@app.get("/bookings", response_model=list[BookingOut])
def bookings_for_day(
    day: date,
    _: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    rows = db.execute(
        select(SpaceBooking, Space.name, User.full_name)
        .join(Space, SpaceBooking.space_id == Space.id)
        .join(User, SpaceBooking.user_id == User.id)
        .where(
            SpaceBooking.booking_date == day,
            SpaceBooking.status == "confirmed",
        )
        .order_by(SpaceBooking.start_minutes)
    ).all()

    return [booking_to_out(b, space_name, user_name) for b, space_name, user_name in rows]


@app.get("/bookings/mine", response_model=list[BookingOut])
def my_bookings(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    rows = db.execute(
        select(SpaceBooking, Space.name, User.full_name)
        .join(Space, SpaceBooking.space_id == Space.id)
        .join(User, SpaceBooking.user_id == User.id)
        .where(
            SpaceBooking.user_id == user.id,
            SpaceBooking.booking_date >= get_work_date() - timedelta(days=30),
        )
        .order_by(SpaceBooking.booking_date.desc(), SpaceBooking.start_minutes.desc())
    ).all()

    return [booking_to_out(b, space_name, user_name) for b, space_name, user_name in rows]


@app.post("/bookings", response_model=BookingOut)
def create_booking(
    data: BookingIn,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    booking, space = book_space(
        db,
        user,
        data.space_id,
        data.booking_date,
        to_minutes(data.start_time),
        to_minutes(data.end_time),
        data.title.strip(),
        (data.notes or "").strip() or None,
    )

    db.commit()
    db.refresh(booking)

    return booking_to_out(booking, space.name, user.full_name)


@app.post("/bookings/{booking_id}/cancel", response_model=BookingOut)
def cancel_booking(
    booking_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    booking = db.get(SpaceBooking, booking_id)

    if not booking:
        raise HTTPException(status_code=404, detail="Reserva no encontrada.")

    if booking.user_id != user.id and user.role != "admin":
        raise HTTPException(status_code=403, detail="No puedes cancelar esta reserva.")

    if booking.status != "confirmed":
        raise HTTPException(status_code=409, detail="Esta reserva ya fue cancelada.")

    if booking.booking_date < get_work_date():
        raise HTTPException(status_code=409, detail="No se pueden cancelar reservas pasadas.")

    booking.status = "cancelled"

    space = db.get(Space, booking.space_id)
    owner = db.get(User, booking.user_id)

    if booking.user_id != user.id:
        create_notification(
            db,
            booking.user_id,
            "booking",
            "Reserva cancelada",
            (
                f"Se canceló tu reserva de {space.name} el "
                f"{booking.booking_date:%d/%m/%Y} a las {to_hhmm(booking.start_minutes)}."
            ),
            "spaces",
            f"booking-{booking.id}-cancelled",
        )

    db.commit()
    db.refresh(booking)

    return booking_to_out(booking, space.name, owner.full_name)

def purge_user(db: Session, user: User) -> None:
    uid = user.id

    # Lo que este usuario hizo sobre otros se conserva, sin su nombre
    db.execute(update(EmployeeProfile).where(EmployeeProfile.manager_id == uid).values(manager_id=None))
    db.execute(update(AttendanceRecord).where(AttendanceRecord.created_by_id == uid).values(created_by_id=None))
    db.execute(update(VacationRequest).where(VacationRequest.reviewed_by_id == uid).values(reviewed_by_id=None))
    db.execute(update(PermissionRequest).where(PermissionRequest.reviewed_by_id == uid).values(reviewed_by_id=None))
    db.execute(delete(QrSession).where(QrSession.created_by_id == uid))

    own_posts = db.scalars(select(Post.id).where(Post.author_id == uid)).all()

    if own_posts:
        db.execute(delete(PostResponse).where(PostResponse.post_id.in_(own_posts)))
        db.execute(delete(Post).where(Post.id.in_(own_posts)))

    db.execute(delete(PostResponse).where(PostResponse.user_id == uid))

    # Los datos propios del usuario se eliminan
    for model in (
        AttendanceRecord,
        VacationRequest,
        PermissionRequest,
        Notification,
        UserPreferences,
        SpaceBooking,
    ):
        db.execute(delete(model).where(model.user_id == uid))

    db.delete(user)


@app.delete("/users/{user_id}")
def delete_user(
    user_id: int,
    confirm: str,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    target = db.get(User, user_id)

    if not target:
        raise HTTPException(status_code=404, detail="Usuario no encontrado.")

    if target.id == admin.id:
        raise HTTPException(status_code=400, detail="No puedes eliminar tu propia cuenta.")

    if confirm.strip().lower() != target.email.lower():
        raise HTTPException(status_code=400, detail="El correo de confirmación no coincide.")

    purge_user(db, target)
    db.commit()

    return {"ok": True}

def nth_weekday(year: int, month: int, weekday: int, n: int) -> date:
    """n-ésimo día de la semana (lunes = 0) de un mes."""

    first = date(year, month, 1)
    return first + timedelta(days=(weekday - first.weekday()) % 7 + 7 * (n - 1))


def statutory_holidays(year: int) -> list[tuple[date, str]]:
    items = [
        (date(year, 1, 1), "Año Nuevo"),
        (nth_weekday(year, 2, 0, 1), "Día de la Constitución"),
        (nth_weekday(year, 3, 0, 3), "Natalicio de Benito Juárez"),
        (date(year, 5, 1), "Día del Trabajo"),
        (date(year, 9, 16), "Día de la Independencia"),
        (nth_weekday(year, 11, 0, 3), "Día de la Revolución"),
        (date(year, 12, 25), "Navidad"),
    ]

    if (year - 1994) % 6 == 0:
        items.append((date(year, 10, 1), "Transmisión del Poder Ejecutivo Federal"))

    return sorted(items)


@app.get("/holidays", response_model=list[HolidayOut])
def list_holidays(
    _: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    return db.scalars(select(Holiday).order_by(Holiday.holiday_date)).all()


@app.post("/holidays", response_model=HolidayOut)
def create_holiday(
    data: HolidayIn,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    if db.scalar(select(Holiday.id).where(Holiday.holiday_date == data.holiday_date)):
        raise HTTPException(status_code=409, detail="Ya existe un festivo en esa fecha.")

    holiday = Holiday(holiday_date=data.holiday_date, name=data.name.strip())
    db.add(holiday)
    db.commit()
    db.refresh(holiday)

    return holiday


@app.post("/holidays/statutory")
def import_statutory_holidays(
    year: int,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    if not 2000 <= year <= 2100:
        raise HTTPException(status_code=400, detail="Año fuera de rango.")

    existing = set(
        db.scalars(
            select(Holiday.holiday_date).where(
                Holiday.holiday_date >= date(year, 1, 1),
                Holiday.holiday_date <= date(year, 12, 31),
            )
        ).all()
    )

    created = 0

    for day, name in statutory_holidays(year):
        if day not in existing:
            db.add(Holiday(holiday_date=day, name=name))
            created += 1

    db.commit()

    return {"created": created}


@app.delete("/holidays/{holiday_id}")
def delete_holiday(
    holiday_id: int,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    holiday = db.get(Holiday, holiday_id)

    if not holiday:
        raise HTTPException(status_code=404, detail="Festivo no encontrado.")

    db.delete(holiday)
    db.commit()

    return {"ok": True}

# =========================================================
# PERSONAL
# =========================================================


def schedule_out(schedule: Schedule) -> ScheduleOut:
    return ScheduleOut(
        start_time=to_hhmm(schedule.start),
        end_time=to_hhmm(schedule.end),
        tolerance_minutes=schedule.tolerance,
        work_days=schedule.days,
        is_custom=schedule.custom,
    )


def build_person(user: User, profile: "EmployeeProfile | None", phone: str | None) -> PersonOut:
    return PersonOut(
        id=user.id,
        full_name=user.full_name,
        email=user.email,
        role=user.role,
        avatar=user.avatar,
        is_active=user.is_active,
        hire_date=user.hire_date,
        employee_code=profile.employee_code if profile else None,
        position=profile.position if profile else None,
        department=profile.department if profile else None,
        manager_id=profile.manager_id if profile else None,
        phone=phone,
        is_top_level=bool(profile.is_top_level) if profile else False,
        schedule=schedule_out(schedule_from_profile(profile)),
    )


def creates_cycle(db: Session, user_id: int, manager_id: int) -> bool:
    """¿Asignar este jefe haría que la persona termine reportándose a sí misma?"""

    managers = {p.user_id: p.manager_id for p in db.scalars(select(EmployeeProfile)).all()}
    current, seen = manager_id, set()

    while current is not None and current not in seen:
        if current == user_id:
            return True

        seen.add(current)
        current = managers.get(current)

    return False


@app.get("/people", response_model=list[PersonOut])
def list_people(
    include_inactive: bool = False,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    is_admin = user.role == "admin"
    query = select(User).order_by(User.full_name)

    if not (is_admin and include_inactive):
        query = query.where(User.is_active.is_(True))

    profiles = {p.user_id: p for p in db.scalars(select(EmployeeProfile)).all()}

    # El teléfono solo lo ve el admin
    phones = (
        {p.user_id: p.phone for p in db.scalars(select(UserPreferences)).all()}
        if is_admin
        else {}
    )

    return [
        build_person(person, profiles.get(person.id), phones.get(person.id))
        for person in db.scalars(query).all()
    ]


@app.get("/schedule/mine", response_model=ScheduleOut)
def my_schedule(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    return schedule_out(get_schedule(db, user.id))


@app.put("/people/{user_id}/profile", response_model=PersonOut)
def update_person_profile(
    user_id: int,
    data: ProfileUpdateIn,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    manager_id = None if data.is_top_level else data.manager_id
    target = db.get(User, user_id)

    if not target:
        raise HTTPException(status_code=404, detail="Persona no encontrada.")

    code = (data.employee_code or "").strip() or None

    if code and db.scalar(
        select(EmployeeProfile.id).where(
            EmployeeProfile.employee_code == code,
            EmployeeProfile.user_id != user_id,
        )
    ):
        raise HTTPException(status_code=409, detail="Ese código de empleado ya está asignado a otra persona.")

    if manager_id is not None:
        if manager_id == user_id:
            raise HTTPException(status_code=400, detail="Una persona no puede ser su propio jefe.")

        if not db.get(User, manager_id):
            raise HTTPException(status_code=404, detail="El jefe directo no existe.")

        if creates_cycle(db, user_id, manager_id):
            raise HTTPException(
                status_code=400,
                detail="Ese jefe depende de esta persona: se formaría un ciclo en el organigrama.",
            )

    profile = get_profile(db, user_id)

    if profile is None:
        profile = EmployeeProfile(user_id=user_id)
        db.add(profile)

    profile.employee_code = code
    profile.position = (data.position or "").strip() or None
    profile.department = (data.department or "").strip() or None
    profile.manager_id = manager_id
    target.hire_date = data.hire_date

    if data.custom_schedule:
        profile.work_start_minutes = to_minutes(data.start_time)
        profile.work_end_minutes = to_minutes(data.end_time)
        profile.tolerance_minutes = data.tolerance_minutes
        profile.work_days = ",".join(str(d) for d in sorted(set(data.work_days)))
    else:
        profile.work_start_minutes = None
        profile.work_end_minutes = None
        profile.tolerance_minutes = None
        profile.work_days = None

    profile.is_top_level = data.is_top_level

    db.commit()
    db.refresh(target)
    db.refresh(profile)

    return build_person(target, profile, None)

# =========================================================
# CARGA MASIVA DESDE EXCEL
# =========================================================

IMPORT_COLUMNS = [
    "Código de empleado",
    "Nombre completo",
    "Correo",
    "Rol",
    "Puesto",
    "Departamento",
    "Correo del jefe directo",
    "Fecha de ingreso",
    "Teléfono",
    "Hora de entrada",
    "Hora de salida",
    "Tolerancia (min)",
    "Días laborables",
]

IMPORT_EXAMPLE = [
    "E-001", "Ana López Pérez", "ana.lopez@empresa.com", "Empleado", "Contadora",
    "Contabilidad", "director@empresa.com", "2024-03-15", "+52 999 123 4567",
    "09:00", "18:00", "10", "L,M,X,J,V",
]

IMPORT_REQUIRED = {"Nombre completo", "Correo"}

IMPORT_GUIDE = [
    "CÓMO LLENAR LA PLANTILLA DE PERSONAL",
    "",
    "• Una fila por persona desde la fila 2. Borra la fila de ejemplo (en gris) antes de importar.",
    "• Obligatorias (*): Nombre completo y Correo.",
    "• Si el correo ya existe, la persona se ACTUALIZA. Si no, se CREA con una contraseña temporal.",
    "• Las celdas vacías NO borran datos que ya existan en el sistema.",
    "• Rol: Empleado o Administrador. Si lo dejas vacío, será Empleado.",
    "• Correo del jefe directo: puede estar en este mismo archivo o ya existir en el sistema.",
    "• Fecha de ingreso: AAAA-MM-DD (ej. 2024-03-15) o DD/MM/AAAA. Define antigüedad y días de vacaciones.",
    "• Código de empleado: el mismo que usas en nómina. Debe ser único.",
    "• Horario propio (opcional): Hora de entrada y de salida en formato 24 h (09:00), tolerancia en minutos",
    "  y días laborables con letras separadas por comas: L, M, X, J, V, S, D. Si lo dejas vacío, se usa el general.",
    "• Máximo 500 filas por archivo. El sistema revisa todo antes de guardar: si hay errores, no se guarda nada.",
]

DAY_KEYS = {"L": 0, "M": 1, "X": 2, "J": 3, "V": 4, "S": 5, "D": 6}
ROLE_KEYS = {"": "employee", "empleado": "employee", "administrador": "admin", "admin": "admin"}
EMAIL_ADAPTER = TypeAdapter(EmailStr)


def cell_text(value) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def parse_date_cell(value) -> date | None:
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    raise ValueError("Fecha de ingreso inválida (usa AAAA-MM-DD).")


def parse_time_cell(value) -> int | None:
    if value in (None, ""):
        return None

    if isinstance(value, (time, datetime)):
        return value.hour * 60 + value.minute

    if isinstance(value, float):  # Excel guarda horas como fracción del día
        return round(value * 24 * 60)

    text = str(value).strip()[:5]

    if not re.match(HHMM, text):
        raise ValueError(f"Hora '{value}' inválida (usa HH:MM en 24 h).")

    return to_minutes(text)


def parse_days_cell(value) -> list[int] | None:
    text = cell_text(value).upper()

    if not text:
        return None

    days = []

    for part in re.split(r"[,\s;]+", text):
        if not part:
            continue

        if part not in DAY_KEYS:
            raise ValueError(f"Día '{part}' no válido (usa L, M, X, J, V, S, D).")

        days.append(DAY_KEYS[part])

    return sorted(set(days))


@app.get("/people/import/template")
def import_template(_: User = Depends(admin_user)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Personal"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E4BA0")

    for column, title in enumerate(IMPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=f"{title} *" if title in IMPORT_REQUIRED else title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

        example = sheet.cell(row=2, column=column, value=IMPORT_EXAMPLE[column - 1])
        example.font = Font(italic=True, color="8A94A6")

        sheet.column_dimensions[cell.column_letter].width = max(16, len(title) + 6)

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"

    roles = DataValidation(type="list", formula1='"Empleado,Administrador"', allow_blank=True)
    sheet.add_data_validation(roles)
    roles.add("D2:D501")

    guide = workbook.create_sheet("Instrucciones")

    for row, line in enumerate(IMPORT_GUIDE, start=1):
        guide.cell(row=row, column=1, value=line).font = Font(bold=row == 1, size=13 if row == 1 else 11)

    guide.column_dimensions["A"].width = 120

    buffer = BytesIO()
    workbook.save(buffer)

    return {
        "filename": "plantilla-personal.xlsx",
        "content": base64.b64encode(buffer.getvalue()).decode(),
    }


@app.post("/people/import", response_model=ImportResult)
def import_people(
    data: ImportIn,
    admin: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    try:
        raw = base64.b64decode(data.content, validate=True)
    except Exception:
        raise HTTPException(status_code=400, detail="El archivo no se pudo leer.")

    if len(raw) > 2_000_000:
        raise HTTPException(status_code=400, detail="El archivo supera 2 MB.")

    try:
        workbook = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo abrir el archivo. Usa la plantilla en formato .xlsx.")

    sheet = workbook["Personal"] if "Personal" in workbook.sheetnames else workbook.worksheets[0]

    users_by_email = {u.email.lower(): u for u in db.scalars(select(User)).all()}
    users_by_id = {u.id: u for u in users_by_email.values()}
    profiles = {p.user_id: p for p in db.scalars(select(EmployeeProfile)).all()}
    code_owner = {p.employee_code: p.user_id for p in profiles.values() if p.employee_code}

    parsed: list[dict] = []
    seen_emails: set[str] = set()
    seen_codes: set[str] = set()
    width = len(IMPORT_COLUMNS)

    # ---------- 1. Leer y validar cada fila ----------
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(values[:width]) + [None] * (width - len(values[:width]))
        text = [cell_text(v) for v in cells]

        if not any(text):
            continue

        email = text[2].lower()

        if email == IMPORT_EXAMPLE[2]:
            continue  # fila de ejemplo

        if len(parsed) >= 500:
            raise HTTPException(status_code=400, detail="Máximo 500 filas por archivo.")

        errors: list[str] = []
        existing = users_by_email.get(email)

        if len(text[1]) < 2:
            errors.append("Falta el nombre completo.")

        try:
            EMAIL_ADAPTER.validate_python(email)
        except Exception:
            errors.append("Correo inválido o vacío.")

        if email in seen_emails:
            errors.append("Correo repetido en el archivo.")
        seen_emails.add(email)

        role = ROLE_KEYS.get(text[3].lower())
        if role is None:
            errors.append("Rol inválido: usa Empleado o Administrador.")

        code = text[0] or None
        if code:
            if code in seen_codes:
                errors.append("Código de empleado repetido en el archivo.")
            seen_codes.add(code)

            owner = code_owner.get(code)
            if owner and (not existing or owner != existing.id):
                errors.append("Ese código ya pertenece a otra persona.")

        hire = start = end = tolerance = days = None

        try:
            hire = parse_date_cell(cells[7])
        except ValueError as exc:
            errors.append(str(exc))

        try:
            start = parse_time_cell(cells[9])
            end = parse_time_cell(cells[10])
        except ValueError as exc:
            errors.append(str(exc))

        if text[11]:
            if text[11].isdigit() and int(text[11]) <= 180:
                tolerance = int(text[11])
            else:
                errors.append("Tolerancia inválida (0 a 180 minutos).")

        try:
            days = parse_days_cell(cells[12])
        except ValueError as exc:
            errors.append(str(exc))

        custom = any(v is not None for v in (start, end, tolerance, days))

        if custom and (start is None or end is None):
            errors.append("Para un horario propio indica hora de entrada y de salida.")

        if start is not None and end is not None and end <= start:
            errors.append("La salida debe ser posterior a la entrada.")

        manager = text[6].lower() or None

        if manager and manager == email:
            errors.append("Una persona no puede ser su propio jefe.")

        parsed.append({
            "row": row_number, "email": email, "name": text[1], "role": role, "code": code,
            "position": text[4] or None, "department": text[5] or None, "manager": manager,
            "hire": hire, "phone": text[8] or None, "start": start, "end": end,
            "tolerance": tolerance, "days": days, "custom": custom,
            "errors": errors, "existing": existing,
        })

    # ---------- 2. Jefes: que existan y que no formen ciclos ----------
    file_emails = {p["email"] for p in parsed}
    chain: dict[str, str] = {}

    for profile in profiles.values():
        member = users_by_id.get(profile.user_id)
        boss = users_by_id.get(profile.manager_id) if profile.manager_id else None

        if member and boss:
            chain[member.email.lower()] = boss.email.lower()

    for item in parsed:
        if not item["manager"]:
            continue

        if item["manager"] not in file_emails and item["manager"] not in users_by_email:
            item["errors"].append("El jefe directo no existe ni en el archivo ni en el sistema.")
        else:
            chain[item["email"]] = item["manager"]

    for item in parsed:
        current, seen = chain.get(item["email"]), set()

        while current and current not in seen:
            if current == item["email"]:
                item["errors"].append("Los jefes forman un ciclo (A reporta a B y B reporta a A).")
                break

            seen.add(current)
            current = chain.get(current)

    rows = [
        ImportRowResult(
            row=p["row"],
            email=p["email"],
            full_name=p["name"],
            action="error" if p["errors"] else ("update" if p["existing"] else "create"),
            errors=p["errors"],
        )
        for p in parsed
    ]

    summary = {
        "rows": rows,
        "created": sum(r.action == "create" for r in rows),
        "updated": sum(r.action == "update" for r in rows),
        "failed": sum(r.action == "error" for r in rows),
    }

    if data.dry_run:
        return ImportResult(dry_run=True, **summary)

    if not parsed:
        raise HTTPException(status_code=400, detail="El archivo no tiene filas para importar.")

    if summary["failed"]:
        raise HTTPException(status_code=400, detail="Corrige los errores del archivo antes de importar.")

    # ---------- 3. Guardar todo o nada ----------
    credentials: list[ImportCredential] = []

    try:
        for item in parsed:
            user = item["existing"]

            if user is None:
                temporary = generate_temporary_password()
                user = User(
                    email=item["email"],
                    full_name=item["name"],
                    password_hash=password_hash.hash(temporary),
                    role=item["role"],
                    is_active=True,
                )
                db.add(user)
                db.flush()

                users_by_email[item["email"]] = user
                credentials.append(
                    ImportCredential(full_name=item["name"], email=item["email"], temporary_password=temporary)
                )
            else:
                user.full_name = item["name"]

                if user.id != admin.id:  # nadie se quita su propio rol de admin
                    user.role = item["role"]

            if item["hire"]:
                user.hire_date = item["hire"]

            profile = profiles.get(user.id)

            if profile is None:
                profile = EmployeeProfile(user_id=user.id)
                db.add(profile)
                profiles[user.id] = profile

            if item["code"]:
                profile.employee_code = item["code"]
            if item["position"]:
                profile.position = item["position"]
            if item["department"]:
                profile.department = item["department"]

            if item["custom"]:
                profile.work_start_minutes = item["start"]
                profile.work_end_minutes = item["end"]
                profile.tolerance_minutes = (
                    item["tolerance"] if item["tolerance"] is not None else settings.late_grace_minutes
                )
                profile.work_days = ",".join(str(d) for d in (item["days"] or [0, 1, 2, 3, 4]))

            if item["phone"]:
                prefs = db.scalar(select(UserPreferences).where(UserPreferences.user_id == user.id))

                if prefs is None:
                    prefs = UserPreferences(user_id=user.id)
                    db.add(prefs)

                prefs.phone = item["phone"]

        db.flush()

        # Los jefes se asignan al final: pueden venir en el mismo archivo
        for item in parsed:
            if item["manager"]:
                member = users_by_email[item["email"]]
                profiles[member.id].manager_id = users_by_email[item["manager"]].id

        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="No se pudo completar la importación. No se guardó ningún cambio.",
        ) from exc

    return ImportResult(dry_run=False, **summary, credentials=credentials)

@app.get("/reports/codes", response_model=list[IncidenceCodeOut])
def list_incidence_codes(
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    return load_codes(db)


@app.put("/reports/codes", response_model=list[IncidenceCodeOut])
def update_incidence_codes(
    data: list[IncidenceCodeIn],
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    if {item.key for item in data} - set(REPORT_KEYS):
        raise HTTPException(status_code=400, detail="Hay una clave interna desconocida.")

    codes = [item.code for item in data]

    if len(codes) != len(set(codes)):
        raise HTTPException(status_code=400, detail="Cada clave debe ser única.")

    rows = {c.key: c for c in db.scalars(select(IncidenceCode)).all()}

    for item in data:
        row = rows.get(item.key)

        if row is None:
            row = IncidenceCode(key=item.key)
            db.add(row)

        row.code = item.code
        row.label = item.label.strip()
        row.include_in_incidences = item.include_in_incidences

    db.commit()

    return load_codes(db)


@app.post("/reports/attendance/preview", response_model=ReportPreviewOut)
def report_preview(
    data: ReportRequest,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    users, profiles, days = build_report_days(db, data)
    stats = summarize_report(users, days)

    people = []

    for user in users:
        code, department, position = employee_fields(user, profiles)
        s = stats[user.id]

        people.append(
            ReportPersonOut(
                user_id=user.id,
                employee_code=code or None,
                full_name=user.full_name,
                department=department or None,
                position=position or None,
                workdays=s["workdays"],
                counts=s["counts"],
                hours=round(s["hours"], 2),
                late_minutes=s["late"],
            )
        )

    return ReportPreviewOut(
        start=data.start,
        end=data.end,
        codes=load_codes(db),
        people=people,
        totals={k: sum(p.counts[k] for p in people) for k in REPORT_KEYS},
        hours=round(sum(p.hours for p in people), 2),
        late_minutes=sum(p.late_minutes for p in people),
    )


@app.post("/reports/attendance/export")
def report_export(
    data: ReportRequest,
    _: User = Depends(admin_user),
    db: Session = Depends(get_db),
):
    users, profiles, days = build_report_days(db, data)
    stats = summarize_report(users, days)
    codes = {c.key: c for c in load_codes(db)}
    branding = get_branding(db)
    primary = branding.primary_color.lstrip("#").upper()

    days.sort(key=lambda d: (d.user.full_name, d.day))
    workbook = Workbook()

    # ---------- 1. Resumen ----------
    summary = workbook.active
    summary.title = "Resumen"

    summary["A1"] = branding.company_name
    summary["A1"].font = Font(bold=True, size=16)
    summary["A2"] = "Reporte de asistencia"
    summary["A2"].font = Font(bold=True, size=13, color=primary)

    filters = [f"Periodo: {data.start:%d/%m/%Y} al {data.end:%d/%m/%Y}"]
    if data.department:
        filters.append(f"Departamento: {data.department}")

    summary["A3"] = " · ".join(filters)
    summary["A4"] = f"Generado el {get_local_now():%d/%m/%Y a las %H:%M}"

    for ref in ("A3", "A4"):
        summary[ref].font = Font(color="65718A")

    headers = [
        "Código", "Empleado", "Departamento", "Puesto", "Días laborables",
        *[codes[k].code for k in REPORT_KEYS],
        "Horas trabajadas", "Minutos de retardo",
    ]

    top = 6

    for col, title in enumerate(headers, start=1):
        style_header(summary.cell(row=top, column=col, value=title), primary)

    # Las columnas de claves llevan su color y su descripción como nota
    for i, key in enumerate(REPORT_KEYS):
        cell = summary.cell(row=top, column=6 + i)
        cell.fill = solid(CODE_COLORS[key])
        cell.font = Font(bold=True, color="172033")
        cell.comment = Comment(codes[key].label, "Sistema")

    row = top + 1
    totals = [0.0] * (len(headers) - 4)

    for index, user in enumerate(users):
        code, department, position = employee_fields(user, profiles)
        s = stats[user.id]

        numbers = [s["workdays"], *[s["counts"][k] for k in REPORT_KEYS], round(s["hours"], 2), s["late"]]
        totals = [a + b for a, b in zip(totals, numbers)]

        for col, value in enumerate([code, user.full_name, department, position, *numbers], start=1):
            cell = summary.cell(row=row, column=col, value=value)
            cell.border = BORDER

            if index % 2:
                cell.fill = solid("F7F9FC")
            if col >= 5:
                cell.alignment = Alignment(horizontal="center")

        summary.cell(row=row, column=len(headers) - 1).number_format = "0.00"
        row += 1

    summary.cell(row=row, column=2, value="Totales").font = Font(bold=True)

    for offset, value in enumerate(totals):
        cell = summary.cell(row=row, column=5 + offset, value=round(value, 2))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(top=Side(style="medium", color=primary))

    summary.cell(row=row, column=len(headers) - 1).number_format = "0.00"

    for col, width in enumerate([12, 30, 20, 22, 11] + [8] * len(REPORT_KEYS) + [12, 12], start=1):
        summary.column_dimensions[get_column_letter(col)].width = width

    summary.row_dimensions[top].height = 32
    summary.freeze_panes = summary.cell(row=top + 1, column=3)
    summary.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{max(row - 1, top)}"
    summary.page_setup.orientation = "landscape"
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0

    # ---------- 2. Incidencias (lo que se captura en nómina) ----------
    incidents = workbook.create_sheet("Incidencias")
    incident_rows = []

    for item in days:
        code, _, _ = employee_fields(item.user, profiles)
        info = codes[item.key]

        if info.include_in_incidences and item.key != "incomplete":
            incident_rows.append([
                code, item.user.full_name, item.day, info.code, info.label,
                1, item.late_minutes or None, item.note or "",
            ])

        if item.incomplete and codes["incomplete"].include_in_incidences:
            incident_rows.append([
                code, item.user.full_name, item.day, codes["incomplete"].code,
                codes["incomplete"].label, 1, None, "Entrada sin salida registrada",
            ])

    write_table(
        incidents,
        ["Código empleado", "Empleado", "Fecha", "Clave", "Descripción", "Cantidad", "Minutos", "Observaciones"],
        incident_rows,
        primary,
        [16, 30, 12, 9, 32, 10, 10, 40],
    )

    for r in range(2, len(incident_rows) + 2):
        incidents.cell(row=r, column=3).number_format = "DD/MM/YYYY"
        incidents.cell(row=r, column=4).font = Font(bold=True)

    # ---------- 3. Detalle diario ----------
    detail = workbook.create_sheet("Detalle diario")
    detail_rows = []

    for item in days:
        code, department, _ = employee_fields(item.user, profiles)
        info = codes[item.key]

        detail_rows.append([
            code, item.user.full_name, department, item.day, WEEKDAYS[item.day.weekday()],
            info.code, info.label, local_hhmm(item.entry), local_hhmm(item.exit),
            item.hours or None, item.late_minutes or None,
            "Sí" if item.incomplete else "", item.note or "",
        ])

    write_table(
        detail,
        ["Código", "Empleado", "Departamento", "Fecha", "Día", "Clave", "Descripción",
         "Entrada", "Salida", "Horas", "Retardo (min)", "Sin salida", "Observaciones"],
        detail_rows,
        primary,
        [12, 30, 20, 12, 11, 8, 30, 9, 9, 8, 12, 10, 40],
    )

    for r, item in enumerate(days, start=2):
        detail.cell(row=r, column=4).number_format = "DD/MM/YYYY"
        detail.cell(row=r, column=10).number_format = "0.00"

        key_cell = detail.cell(row=r, column=6)
        key_cell.fill = solid(CODE_COLORS[item.key])
        key_cell.font = Font(bold=True)
        key_cell.alignment = Alignment(horizontal="center")

    # ---------- 4. Calendario ----------
    calendar = workbook.create_sheet("Calendario")
    dates = [data.start + timedelta(days=i) for i in range((data.end - data.start).days + 1)]
    by_day = {(item.user.id, item.day): item for item in days}

    for col, title in ((1, "Código"), (2, "Empleado")):
        style_header(calendar.cell(row=1, column=col, value=title), primary)
        style_header(calendar.cell(row=2, column=col), primary)

    for i, day in enumerate(dates):
        col = 3 + i
        style_header(calendar.cell(row=1, column=col, value=day.strftime("%d/%m")), primary)
        style_header(calendar.cell(row=2, column=col, value="LMXJVSD"[day.weekday()]), primary)
        calendar.column_dimensions[get_column_letter(col)].width = 6.5

    for r, user in enumerate(users, start=3):
        code, _, _ = employee_fields(user, profiles)
        calendar.cell(row=r, column=1, value=code).border = BORDER
        calendar.cell(row=r, column=2, value=user.full_name).border = BORDER

        for i, day in enumerate(dates):
            item = by_day.get((user.id, day))
            cell = calendar.cell(row=r, column=3 + i, value=codes[item.key].code if item else "")
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER

            if item:
                cell.fill = solid(CODE_COLORS[item.key])
                cell.font = Font(bold=True, color="B42318" if item.incomplete else "172033")

                if item.incomplete:
                    cell.comment = Comment("Entrada sin salida registrada", "Sistema")

    legend = len(users) + 5
    calendar.cell(row=legend, column=2, value="Leyenda").font = Font(bold=True)

    for i, key in enumerate(REPORT_KEYS, start=1):
        swatch = calendar.cell(row=legend + i, column=1, value=codes[key].code)
        swatch.fill = solid(CODE_COLORS[key])
        swatch.font = Font(bold=True)
        swatch.alignment = Alignment(horizontal="center")
        calendar.cell(row=legend + i, column=2, value=codes[key].label)

    calendar.column_dimensions["A"].width = 12
    calendar.column_dimensions["B"].width = 30
    calendar.freeze_panes = "C3"

    # ---------- 5. Claves ----------
    legend_sheet = workbook.create_sheet("Claves")
    all_totals = {k: sum(stats[u.id]["counts"][k] for u in users) for k in REPORT_KEYS}

    write_table(
        legend_sheet,
        ["Clave", "Descripción", "Se reporta como incidencia", "Total en el periodo"],
        [
            [codes[k].code, codes[k].label, "Sí" if codes[k].include_in_incidences else "No", all_totals[k]]
            for k in REPORT_KEYS
        ],
        primary,
        [10, 36, 26, 20],
    )

    for r, key in enumerate(REPORT_KEYS, start=2):
        legend_sheet.cell(row=r, column=1).fill = solid(CODE_COLORS[key])
        legend_sheet.cell(row=r, column=1).font = Font(bold=True)

    legend_sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)

    return {
        "filename": f"asistencia_{data.start:%Y%m%d}_{data.end:%Y%m%d}.xlsx",
        "content": base64.b64encode(buffer.getvalue()).decode(),
    }

# =========================================================
# COMUNICADOS Y REUNIONES
# =========================================================


@dataclass
class AudienceContext:
    users: dict[int, User]               # usuarios activos
    departments: dict[int, str | None]
    positions: dict[int, str | None]
    children: dict[int, list[int]]       # jefe -> reportes directos


def audience_context(db: Session) -> AudienceContext:
    users = {u.id: u for u in db.scalars(select(User).where(User.is_active.is_(True))).all()}
    profiles = db.scalars(select(EmployeeProfile)).all()
    children: dict[int, list[int]] = {}

    for profile in profiles:
        if profile.manager_id:
            children.setdefault(profile.manager_id, []).append(profile.user_id)

    return AudienceContext(
        users=users,
        departments={p.user_id: p.department for p in profiles},
        positions={p.user_id: p.position for p in profiles},
        children=children,
    )


def team_of(user_id: int, ctx: AudienceContext) -> set[int]:
    """Todas las personas que dependen de alguien, directa o indirectamente."""

    result: set[int] = set()
    pending = list(ctx.children.get(user_id, []))

    while pending:
        member = pending.pop()

        if member not in result:
            result.add(member)
            pending.extend(ctx.children.get(member, []))

    return result


def audience_members(post: Post, ctx: AudienceContext) -> set[int]:
    active = set(ctx.users)

    if post.audience == "team":
        return (team_of(post.author_id, ctx) | {post.author_id}) & active

    if post.audience.startswith("dept:"):
        name = post.audience[5:]
        return {uid for uid in active if ctx.departments.get(uid) == name} | {post.author_id}

    return active

def ensure_meeting_room_free(db: Session, data: PostIn, current: Post | None = None) -> None:
    """Valida la sala ANTES de crear o modificar la publicación."""

    if data.kind != "meeting" or not data.space_id:
        return

    conflict = space_conflict(
        db,
        data.space_id,
        data.meeting_date,
        to_minutes(data.start_time),
        to_minutes(data.end_time),
        exclude_booking_id=current.booking_id if current else None,
    )

    if conflict:
        space = db.get(Space, data.space_id)

        raise HTTPException(
            status_code=409,
            detail=(
                f"{space.name if space else 'La sala'} ya está ocupada de "
                f"{to_hhmm(conflict.start_minutes)} a {to_hhmm(conflict.end_minutes)} ese día "
                f"({conflict.title}). Elige otro horario u otra sala: la publicación no se creó."
            ),
        )

def audience_label(audience: str) -> str:
    if audience == "team":
        return "Equipo de trabajo"

    if audience.startswith("dept:"):
        return audience[5:]

    return "Toda la empresa"


def can_publish(user: User, ctx: AudienceContext) -> bool:
    return user.role == "admin" or bool(ctx.children.get(user.id))


def can_see(post: Post, user: User, ctx: AudienceContext) -> bool:
    return user.role == "admin" or user.id == post.author_id or user.id in audience_members(post, ctx)


def check_audience(audience: str, user: User, ctx: AudienceContext) -> None:
    has_team = bool(ctx.children.get(user.id))

    if audience == "team" and has_team:
        return

    if user.role == "admin":
        departments = {d for d in ctx.departments.values() if d}

        if audience == "all" or (audience.startswith("dept:") and audience[5:] in departments):
            return
    else:
        own = ctx.departments.get(user.id)

        if own and audience == f"dept:{own}":
            return

    raise HTTPException(status_code=400, detail="No puedes publicar para esa audiencia.")


def cancel_post_booking(db: Session, post: Post) -> None:
    if post.booking_id:
        booking = db.get(SpaceBooking, post.booking_id)

        if booking and booking.status == "confirmed":
            booking.status = "cancelled"

        post.booking_id = None


def apply_post(db: Session, post: Post, data: PostIn, author: User) -> bool:
    """Copia los datos a la publicación. Devuelve True si cambió el horario de la reunión."""

    post.kind = data.kind
    post.title = data.title.strip()
    post.body = data.body.strip()
    post.image = data.image
    post.audience = data.audience
    post.pinned = data.pinned if author.role == "admin" else False

    if data.kind != "meeting":
        cancel_post_booking(db, post)
        post.meeting_date = post.start_minutes = post.end_minutes = None
        post.location = post.meeting_url = None
        post.space_id = None
        return False

    start = to_minutes(data.start_time)
    end = to_minutes(data.end_time)

    schedule_changed = (post.meeting_date, post.start_minutes, post.end_minutes) != (
        data.meeting_date,
        start,
        end,
    )

    needs_new_booking = data.space_id and (schedule_changed or data.space_id != post.space_id or not post.booking_id)

    if needs_new_booking or not data.space_id:
        cancel_post_booking(db, post)

    if needs_new_booking:
        booking, _ = book_space(
            db,
            db.get(User, post.author_id) or author,
            data.space_id,
            data.meeting_date,
            start,
            end,
            data.title.strip(),
            "Reunión publicada en Comunicados",
        )
        post.booking_id = booking.id

    post.meeting_date = data.meeting_date
    post.start_minutes = start
    post.end_minutes = end
    post.space_id = data.space_id
    post.location = (data.location or "").strip() or None
    post.meeting_url = (data.meeting_url or "").strip() or None

    return schedule_changed


def notify_post(db: Session, post: Post, ctx: AudienceContext, title: str, key: str) -> None:
    author = ctx.users.get(post.author_id)
    when = (
        f" · {post.meeting_date:%d/%m} a las {to_hhmm(post.start_minutes)}"
        if post.kind == "meeting" and post.meeting_date
        else ""
    )

    for user_id in audience_members(post, ctx) - {post.author_id}:
        create_notification(
            db,
            user_id,
            "company",
            title,
            f"{author.full_name if author else 'Dirección'}: {post.title}{when}",
            "news",
            key,
        )


def build_posts_out(db: Session, posts: list[Post], user: User, ctx: AudienceContext) -> list[PostOut]:
    if not posts:
        return []

    ids = [p.id for p in posts]
    counts = {pid: {"seen": 0, "yes": 0, "no": 0, "maybe": 0} for pid in ids}
    mine: dict[int, str] = {}

    for response in db.scalars(select(PostResponse).where(PostResponse.post_id.in_(ids))).all():
        counts[response.post_id][response.response] = counts[response.post_id].get(response.response, 0) + 1

        if response.user_id == user.id:
            mine[response.post_id] = response.response

    author_ids = {p.author_id for p in posts}
    authors = {u.id: u for u in db.scalars(select(User).where(User.id.in_(author_ids))).all()}

    space_ids = {p.space_id for p in posts if p.space_id}
    spaces = (
        {s.id: s.name for s in db.scalars(select(Space).where(Space.id.in_(space_ids))).all()}
        if space_ids
        else {}
    )

    result = []

    for post in posts:
        author = authors.get(post.author_id)

        result.append(
            PostOut(
                id=post.id,
                kind=post.kind,
                title=post.title,
                body=post.body,
                image=post.image,
                pinned=post.pinned,
                audience=post.audience,
                audience_label=audience_label(post.audience),
                created_at=post.created_at,
                updated_at=post.updated_at,
                author=PostAuthorOut(
                    id=post.author_id,
                    full_name=author.full_name if author else "Usuario",
                    avatar=author.avatar if author else None,
                    position=ctx.positions.get(post.author_id),
                ),
                meeting_date=post.meeting_date,
                start_time=to_hhmm(post.start_minutes) if post.start_minutes is not None else None,
                end_time=to_hhmm(post.end_minutes) if post.end_minutes is not None else None,
                location=post.location,
                meeting_url=post.meeting_url,
                space_id=post.space_id,
                space_name=spaces.get(post.space_id),
                counts=counts[post.id],
                my_response=mine.get(post.id),
                booking_id=post.booking_id,
                audience_size=len(audience_members(post, ctx) - {post.author_id}),
                can_manage=user.role == "admin" or user.id == post.author_id,
            )
        )

    return result


def sync_meeting_reminders(db: Session, user: User) -> None:
    """Recordatorio 1 hora antes, para quien confirmó (asistiré o tal vez)."""

    now_local = get_local_now()
    now_minutes = now_local.hour * 60 + now_local.minute

    meetings = db.scalars(
        select(Post)
        .join(PostResponse, PostResponse.post_id == Post.id)
        .where(
            PostResponse.user_id == user.id,
            PostResponse.response.in_(["yes", "maybe"]),
            Post.kind == "meeting",
            Post.meeting_date == now_local.date(),
            Post.start_minutes > now_minutes,
            Post.start_minutes <= now_minutes + 60,
        )
    ).all()

    for post in meetings:
        create_notification(
            db,
            user.id,
            "meeting_reminder",
            "Tu reunión empieza pronto",
            f"{post.title} · hoy a las {to_hhmm(post.start_minutes)}",
            "news",
            f"reminder-{post.id}",
        )


@app.get("/posts/meta", response_model=PostsMetaOut)
def posts_meta(
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)
    is_admin = user.role == "admin"
    own = ctx.departments.get(user.id)

    return PostsMetaOut(
        can_publish=can_publish(user, ctx),
        can_target_all=is_admin,
        departments=(
            sorted({d for d in ctx.departments.values() if d})
            if is_admin
            else ([own] if own else [])
        ),
        has_team=bool(ctx.children.get(user.id)),
        is_admin=is_admin,
    )


@app.get("/posts", response_model=list[PostOut])
def list_posts(
    kind: str | None = None,
    upcoming: bool = False,
    limit: int = 20,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)
    query = select(Post)

    if upcoming:
        query = query.where(
            Post.kind == "meeting",
            Post.meeting_date >= get_work_date(),
        ).order_by(Post.meeting_date, Post.start_minutes)
    else:
        if kind in ("announcement", "meeting"):
            query = query.where(Post.kind == kind)

        query = query.order_by(Post.pinned.desc(), Post.created_at.desc())

    visible = [p for p in db.scalars(query).all() if can_see(p, user, ctx)]

    return build_posts_out(db, visible[: max(1, min(limit, 100))], user, ctx)


@app.post("/posts", response_model=PostOut)
def create_post(
    data: PostIn,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)

    if not can_publish(user, ctx):
        raise HTTPException(
            status_code=403,
            detail="Solo administradores y personas con equipo a su cargo pueden publicar.",
        )

    check_audience(data.audience, user, ctx)
    ensure_meeting_room_free(db, data)

    if data.image:
        validate_avatar(data.image, 1_000_000)

    post = Post(author_id=user.id)
    apply_post(db, post, data, user)

    db.add(post)
    db.flush()

    notify_post(
        db,
        post,
        ctx,
        "Te invitaron a una reunión" if post.kind == "meeting" else "Nuevo comunicado",
        f"post-{post.id}",
    )

    db.commit()
    db.refresh(post)

    return build_posts_out(db, [post], user, ctx)[0]


@app.put("/posts/{post_id}", response_model=PostOut)
def update_post(
    post_id: int,
    data: PostIn,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)
    post = db.get(Post, post_id)

    if not post:
        raise HTTPException(status_code=404, detail="Publicación no encontrada.")

    if user.role != "admin" and user.id != post.author_id:
        raise HTTPException(status_code=403, detail="No puedes editar esta publicación.")

    check_audience(data.audience, user, ctx)
    ensure_meeting_room_free(db, data, post)

    if data.image and data.image != post.image:
        validate_avatar(data.image, 1_000_000)

    schedule_changed = apply_post(db, post, data, user)

    if post.kind == "meeting" and schedule_changed:
        notify_post(
            db,
            post,
            ctx,
            "Cambió el horario de una reunión",
            f"post-{post.id}-{int(datetime.now(timezone.utc).timestamp())}",
        )

    db.commit()
    db.refresh(post)

    return build_posts_out(db, [post], user, ctx)[0]


@app.delete("/posts/{post_id}")
def delete_post(
    post_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    post = db.get(Post, post_id)

    if not post:
        raise HTTPException(status_code=404, detail="Publicación no encontrada.")

    if user.role != "admin" and user.id != post.author_id:
        raise HTTPException(status_code=403, detail="No puedes eliminar esta publicación.")

    cancel_post_booking(db, post)
    db.execute(delete(PostResponse).where(PostResponse.post_id == post.id))
    db.delete(post)
    db.commit()

    return {"ok": True}


@app.post("/posts/{post_id}/respond", response_model=PostOut)
def respond_post(
    post_id: int,
    data: RespondIn,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)
    post = db.get(Post, post_id)

    if not post or not can_see(post, user, ctx):
        raise HTTPException(status_code=404, detail="Publicación no encontrada.")

    allowed = {"seen"} if post.kind == "announcement" else {"yes", "no", "maybe"}

    if data.response != "clear" and data.response not in allowed:
        raise HTTPException(status_code=400, detail="Respuesta inválida.")

    if post.kind == "meeting" and post.meeting_date and post.meeting_date < get_work_date():
        raise HTTPException(status_code=409, detail="Esta reunión ya pasó.")

    existing = db.scalar(
        select(PostResponse).where(
            PostResponse.post_id == post.id,
            PostResponse.user_id == user.id,
        )
    )

    if data.response == "clear":
        if existing:
            db.delete(existing)
    else:
        if existing is None:
            existing = PostResponse(post_id=post.id, user_id=user.id)
            db.add(existing)

        existing.response = data.response
        existing.comment = (data.comment or "").strip() or None

    db.commit()

    return build_posts_out(db, [post], user, ctx)[0]


@app.get("/posts/{post_id}/responses", response_model=list[ResponsePersonOut])
def post_responses(
    post_id: int,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    ctx = audience_context(db)
    post = db.get(Post, post_id)

    if not post:
        raise HTTPException(status_code=404, detail="Publicación no encontrada.")

    if user.role != "admin" and user.id != post.author_id:
        raise HTTPException(status_code=403, detail="Solo el autor puede ver las respuestas.")

    responses = {
        r.user_id: r
        for r in db.scalars(select(PostResponse).where(PostResponse.post_id == post.id)).all()
    }

    people_ids = (audience_members(post, ctx) - {post.author_id}) | set(responses)
    order = {"yes": 0, "seen": 0, "maybe": 1, "no": 2, None: 3}
    result = []

    for user_id in people_ids:
        person = ctx.users.get(user_id) or db.get(User, user_id)

        if not person:
            continue

        response = responses.get(user_id)

        result.append(
            ResponsePersonOut(
                user_id=user_id,
                full_name=person.full_name,
                avatar=person.avatar,
                position=ctx.positions.get(user_id),
                response=response.response if response else None,
                comment=response.comment if response else None,
            )
        )

    result.sort(key=lambda r: (order.get(r.response, 3), r.full_name))

    return result

# =========================================================
# CALENDARIO
# =========================================================


@app.get("/calendar", response_model=list[CalendarEventOut])
def calendar_events(
    start: date,
    end: date,
    team: bool = False,
    user: User = Depends(authenticated_user),
    db: Session = Depends(get_db),
):
    if end < start or (end - start).days > 120:
        raise HTTPException(status_code=400, detail="Rango de fechas inválido (máximo 120 días).")

    events: list[CalendarEventOut] = []

    # Festivos
    for holiday in db.scalars(
        select(Holiday).where(Holiday.holiday_date >= start, Holiday.holiday_date <= end)
    ).all():
        events.append(CalendarEventOut(
            id=f"holiday-{holiday.id}", type="holiday", title=holiday.name,
            start_date=holiday.holiday_date, end_date=holiday.holiday_date, page="calendar",
        ))

    # Mis vacaciones y permisos (aprobados y pendientes)
    for model, kind, page in ((VacationRequest, "vacation", "vacations"), (PermissionRequest, "permission", "attendance")):
        for item in db.scalars(
            select(model).where(
                model.user_id == user.id,
                model.status.in_(["approved", "pending"]),
                model.start_date <= end,
                model.end_date >= start,
            )
        ).all():
            label = "Vacaciones" if kind == "vacation" else (
                "Permiso con goce" if item.kind == "paid" else "Permiso sin goce"
            )

            events.append(CalendarEventOut(
                id=f"{kind}-{item.id}", type=kind, title=label,
                subtitle="Pendiente de aprobación" if item.status == "pending" else item.reason,
                start_date=max(item.start_date, start), end_date=min(item.end_date, end),
                page=page, pending=item.status == "pending",
            ))

    # Ausencias del equipo (solo admin)
    if team and user.role == "admin":
        names = {u.id: u.full_name for u in db.scalars(select(User)).all()}

        for model, label in ((VacationRequest, "Vacaciones"), (PermissionRequest, "Permiso")):
            for item in db.scalars(
                select(model).where(
                    model.user_id != user.id,
                    model.status == "approved",
                    model.start_date <= end,
                    model.end_date >= start,
                )
            ).all():
                events.append(CalendarEventOut(
                    id=f"team-{label}-{item.id}", type="team_absence",
                    title=f"{names.get(item.user_id, 'Empleado')}: {label}",
                    start_date=max(item.start_date, start), end_date=min(item.end_date, end),
                    page="admin-requests",
                ))

    # Reuniones visibles para mí
    ctx = audience_context(db)
    meeting_bookings: set[int] = set()
    space_names = {s.id: s.name for s in db.scalars(select(Space)).all()}

    for post in db.scalars(
        select(Post).where(
            Post.kind == "meeting",
            Post.meeting_date >= start,
            Post.meeting_date <= end,
        )
    ).all():
        if post.booking_id:
            meeting_bookings.add(post.booking_id)

        if not can_see(post, user, ctx):
            continue

        place = space_names.get(post.space_id) or post.location or ("En línea" if post.meeting_url else None)

        events.append(CalendarEventOut(
            id=f"meeting-{post.id}", type="meeting", title=post.title, subtitle=place,
            start_date=post.meeting_date, end_date=post.meeting_date,
            start_time=to_hhmm(post.start_minutes), end_time=to_hhmm(post.end_minutes),
            page="news",
        ))

    # Mis reservas de sala (sin repetir las de reuniones)
    for booking in db.scalars(
        select(SpaceBooking).where(
            SpaceBooking.user_id == user.id,
            SpaceBooking.status == "confirmed",
            SpaceBooking.booking_date >= start,
            SpaceBooking.booking_date <= end,
        )
    ).all():
        if booking.id in meeting_bookings:
            continue

        events.append(CalendarEventOut(
            id=f"booking-{booking.id}", type="booking", title=booking.title,
            subtitle=space_names.get(booking.space_id),
            start_date=booking.booking_date, end_date=booking.booking_date,
            start_time=to_hhmm(booking.start_minutes), end_time=to_hhmm(booking.end_minutes),
            page="spaces",
        ))

    events.sort(key=lambda e: (e.start_date, e.start_time or ""))

    return events

# =========================================================
# FRONTEND ESTÁTICO
# =========================================================

frontend_path = Path(__file__).resolve().parent / "static"
print(f"[DEBUG] frontend_path = {frontend_path}, exists={frontend_path.exists()}")

if frontend_path.exists():

    assets_path = frontend_path / "assets"

    if assets_path.exists():
        app.mount(
            "/assets",
            StaticFiles(
                directory=assets_path
            ),
            name="assets",
        )

    @app.get("/{full_path:path}")
    async def serve_frontend(
        full_path: str,
    ):

        requested_file = frontend_path / full_path

        if (
            requested_file.exists()
            and requested_file.is_file()
        ):
            return FileResponse(
                requested_file
            )

        index_file = frontend_path / "index.html"

        if index_file.exists():
            return FileResponse(
                index_file
            )

        raise HTTPException(
            status_code=404,
            detail="Página no encontrada.",
        )